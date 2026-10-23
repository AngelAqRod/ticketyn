from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import Spacer
from sqlalchemy import func, select

from ticketyn.models import Node, Responsible
from ticketyn.api.documents import (
    AMBER, BLUE, EMERALD, build_pdf, duration_text, fonts,
    modular_metric_strip, paragraph, report_identification, ticket_pdf_table,
)
from ticketyn.api.ticket_data import RELATED_MODELS, ticket_rows, ticket_statement


def friendly_filters(session, filters):
    values = []
    if filters.status:
        values.append(f"Estado: {'Abierto' if filters.status == 'OPEN' else 'Cerrado'}")
    if filters.search:
        values.append(f'Búsqueda: {filters.search}')
    labels = ['Cliente', 'Circuito', 'Sector', 'Departamento', 'Tipo de incidencia']
    for (field, model), label in zip(RELATED_MODELS.items(), labels):
        id = getattr(filters, field)
        if id is None: continue
        item = session.get(model, id)
        if item is None:
            value = 'No disponible'
        elif field == 'customer_id':
            value = f'{item.customer_code} — {item.name}'
        elif field == 'circuit_id':
            value = f'{item.circuit_code} — {item.description}'
        else:
            value = item.name
        values.append(f'{label}: {value}')
    for field, model, label in [("node_id", Node, "Nodo"), ("responsible_id", Responsible, "Responsable")]:
        if getattr(filters, field):
            item = session.get(model, getattr(filters, field))
            values.append(f"{label}: {item.name if item else 'No disponible'}")
    zone = ZoneInfo(filters.timezone)
    if filters.from_at:
        values.append(f'Desde (incluido): {filters.from_at.astimezone(zone):%d/%m/%Y %H:%M}')
    if filters.to_at:
        # Mostrar el límite real evita perder precisión en filtros API intradía.
        values.append(f'Hasta (excluido): {filters.to_at.astimezone(zone):%d/%m/%Y %H:%M}')
    return values or ['Todos los tickets']


def export_ticket_pdf(session, filters, output):
    fonts()
    zone = ZoneInfo(filters.timezone)
    values = filters.model_dump()
    stmt = ticket_statement(values).order_by(None).subquery()
    counts = session.execute(
        select(
            func.count().label('total'),
            func.count().filter(stmt.c.status == 'OPEN').label('open_count'),
            func.count().filter(stmt.c.status == 'CLOSED').label('closed_count'),
        ).select_from(stmt)
    ).one()
    total = counts.total or 0
    open_count = counts.open_count or 0
    closed_count = counts.closed_count or 0

    generated = datetime.now(timezone.utc).astimezone(zone)
    from_str = filters.from_at.astimezone(zone).strftime('%d/%m/%Y %H:%M') if filters.from_at else None
    to_str = filters.to_at.astimezone(zone).strftime('%d/%m/%Y %H:%M') if filters.to_at else None
    if from_str and to_str:
        period_str = f'{from_str} — {to_str}'
    elif from_str:
        period_str = f'Desde (incluido): {from_str}'
    elif to_str:
        period_str = f'Hasta (excluido): {to_str}'
    else:
        period_str = 'Sin restricción temporal'

    filter_tokens = friendly_filters(session, filters)
    filters_str = ' · '.join(filter_tokens)
    width = landscape(A4)[0] - 60

    metrics = [
        ('TOTAL TICKETS', total, BLUE),
        ('ABIERTOS / EN CURSO', open_count, AMBER),
        ('CERRADOS', closed_count, EMERALD),
    ]

    story = [
        report_identification(
            'Listado de tickets',
            period=period_str,
            filters=filters_str,
            total=total,
            accent=BLUE,
            category='Gestión de Incidencias · Tickets de Clientes',
            width=width,
        ),
        modular_metric_strip(metrics, width, accent=BLUE),
        ticket_pdf_table(ticket_rows(session, values), zone, width, light=True, accent=BLUE),
    ]
    build_pdf(output, 'Listado de tickets', filters.timezone, story,
              issued_at=generated.strftime('%d/%m/%Y %H:%M'),
              subtitle='GESTIÓN DE INCIDENCIAS · CLIENTES',
              modular=True, landscape_mode=True, reference='Listado general', accent=BLUE)


def export_ticket_xlsx(session, filters, output):
    headers = ['Referencia', 'Estado', 'Cliente', 'Código cliente', 'Circuito', 'Código circuito',
        'Sector', 'Departamento', 'Tipo de incidencia', 'Título', 'Descripción', 'Inicio', 'Fin',
        'Duración', 'Duración (segundos)', 'Nodo', 'Responsable']
    widths = [22, 14, 32, 24, 38, 28, 26, 26, 28, 45, 65, 34, 34, 22, 22, 28, 28]
    zone = ZoneInfo(filters.timezone)
    def rows():
        for ticket, customer, circuit, sector, department, kind in ticket_rows(session, filters.model_dump()):
            # ISO incluye offset: conserva incluso las horas repetidas al cambiar DST.
            date = lambda value: value.astimezone(zone).isoformat() if value else None
            yield [ticket.reference, 'Abierto' if ticket.status == 'OPEN' else 'Cerrado', customer.name,
                customer.customer_code, circuit.description, circuit.circuit_code, sector.name, department.name,
                kind.name, ticket.title, ticket.description, date(ticket.start_at), date(ticket.end_at),
                duration_text(ticket.duration_seconds) if ticket.duration_seconds is not None else ('En curso' if ticket.status == 'OPEN' else 'Sin fecha de fin'), ticket.duration_seconds, ticket.node.name if ticket.node else None, ticket.responsible.name if ticket.responsible else None]
    write_export_xlsx(output, 'Tickets', headers, widths, rows())


def write_export_xlsx(output, title, headers, widths, rows, *, metadata=None, extra_sheets=()):
    """Shared write-only workbook, literal text cells and existing Tickets styling."""
    workbook = Workbook(write_only=True)
    def append_sheet(name, labels, sizes, records):
        sheet = workbook.create_sheet(name)
        sheet.freeze_panes = 'A2'
        for index, width in enumerate(sizes, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        cells = []
        for label in labels:
            cell = WriteOnlyCell(sheet, value=label)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='102744')
            cells.append(cell)
        sheet.append(cells)
        count = 1
        for values in records:
            cells = []
            for value in values:
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str): cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cells.append(cell)
            sheet.append(cells)
            count += 1
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(labels))}{count}'
    append_sheet(title, headers, widths, rows)
    for name, labels, sizes, records in extra_sheets:
        append_sheet(name, labels, sizes, records)
    if metadata is not None:
        info = workbook.create_sheet('Filtros')
        info.column_dimensions['A'].width = 110
        for value in metadata:
            cell = WriteOnlyCell(info, value=value)
            cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            info.append([cell])
    workbook.save(output)


def stream_ticket_export(session, format, filters, filename, renderer):
    """Same snapshot, spooling and cleanup for client and internal ticket lists."""
    from tempfile import SpooledTemporaryFile
    from fastapi.responses import StreamingResponse
    from starlette.background import BackgroundTask
    from ticketyn.api.reports import read_snapshot
    read_snapshot(session)
    output = SpooledTemporaryFile(max_size=2 * 1024 * 1024)
    try:
        renderer(session, filters, output)
        output.seek(0)
    except Exception:
        output.close()
        raise
    def chunks():
        try:
            while chunk := output.read(65536): yield chunk
        finally:
            output.close()
    media = 'application/pdf' if format == 'pdf' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return StreamingResponse(chunks(), media_type=media,
        headers={'Content-Disposition': f'attachment; filename="{filename}.{format}"'},
        background=BackgroundTask(output.close))
