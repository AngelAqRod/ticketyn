"""PDF de lectura y XLSX de análisis a partir del mismo motor de reportes."""
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ticketyn.api.documents import duration_text
from ticketyn.api.report_data import report_tickets


def metadata(summary):
    zone = ZoneInfo(summary.period.timezone)
    rows = [
        ['Ticketyn', report_title(summary)],
        ['Sector', summary.sector['name'] if summary.sector else 'Todos'],
        ['Nodo', summary.node['name'] if summary.node else 'Todos'],
        ['Responsable', summary.responsible['name'] if summary.responsible else 'Todos'],
        ['Zona horaria', summary.period.timezone],
        ['Desde (incluido)', summary.period.from_at.astimezone(zone).isoformat() if summary.period.from_at else 'Todos'],
        ['Hasta (excluido)', summary.period.to_exclusive.astimezone(zone).isoformat() if summary.period.to_exclusive else 'Sin límite'],
        ['Generado', summary.generated_at.astimezone(zone).isoformat()],
        ['Incidencias iniciadas', summary.kpis.started], ['Incidencias cerradas', summary.kpis.closed],
        ['Duración promedio (segundos)', summary.kpis.average_duration_seconds],
        ['Duración acumulada (segundos)', summary.kpis.total_duration_seconds],
        ['Universo detalle/rankings', 'Inicio en el período.'],
        ['Rankings Nodo/Responsable', 'Solo tickets con asignación; NULL continúa incluido en los KPI correspondientes.'],
        ['Universo cierres/duración', 'Estado CLOSED y Fin en el período, aunque Inicio sea anterior.'],
    ]

    if getattr(summary, 'report_kind', None) == 'internal':
        rows.extend([['Filtros', ' · '.join(summary.applied_filters)], ['Abiertos (iniciados)', summary.open_tickets],
                     ['Cerrados (iniciados)', summary.closed_tickets], ['Tickets de clientes vinculados', summary.related_tickets],
                     ['Incidencias con tickets vinculados', summary.with_related], ['Incidencias sin tickets vinculados', summary.without_related]])
    return rows


def export_xlsx(session, filters, summary, output):
    workbook = Workbook(write_only=True)
    zone = ZoneInfo(filters.timezone)

    def sheet(name, headers, rows, widths):
        page = workbook.create_sheet(name)
        page.freeze_panes = 'A2'
        for index, width in enumerate(widths, 1):
            page.column_dimensions[get_column_letter(index)].width = width
        header = []
        for value in headers:
            cell = WriteOnlyCell(page, value=value)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='166534')
            header.append(cell)
        page.append(header)
        count = 1
        for row in rows:
            cells = []
            for value in row:
                cell = WriteOnlyCell(page, value=value)
                if isinstance(value, str):
                    # Los datos introducidos por el usuario nunca se convierten en fórmulas.
                    cell.data_type = 's'
                elif isinstance(value, datetime):
                    cell.number_format = 'yyyy-mm-dd hh:mm:ss'
                cells.append(cell)
            page.append(cells)
            count += 1
        page.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{count}'
    sheet('Resumen', ['Concepto', 'Valor'], metadata(summary), [38, 95])
    def ticket_rows():
        for ticket, customer, circuit, sector, department, incident_type in report_tickets(session, filters):
            local = lambda value: value.astimezone(zone).replace(tzinfo=None) if value else None
            yield [ticket.reference, ticket.status, customer.customer_code, customer.name,
                   circuit.circuit_code, circuit.description, sector.name, department.name,
                   incident_type.name, ticket.title, ticket.description, local(ticket.start_at),
                   local(ticket.end_at), ticket.duration_seconds, (duration_text(ticket.duration_seconds) if ticket.duration_seconds is not None else ('En curso' if ticket.status == 'OPEN' else 'Sin fecha de fin')),
                   ticket.start_at.astimezone(zone).isoformat(), ticket.end_at.astimezone(zone).isoformat() if ticket.end_at else None, ticket.node.name if ticket.node else None, ticket.responsible.name if ticket.responsible else None]
    if getattr(summary, 'report_kind', None) == 'internal':
        from ticketyn.api.internal_ticket_exports import HEADERS, export_rows
        sheet('Tickets internos', HEADERS, export_rows(session, filters, excel=True), [24, 50, 14, 26, 26, 26, 28, 32, 30, 34, 22, 24])
    else:
        sheet('Tickets', ['Referencia', 'Estado', 'Cliente código', 'Cliente nombre', 'Circuito código',
                         'Circuito descripción', 'Sector', 'Departamento', 'Tipo de incidencia', 'Título',
                         'Descripción', 'Inicio', 'Fin', 'Duración (segundos)', 'Duración', 'Inicio ISO (con offset)', 'Fin ISO (con offset)', 'Nodo', 'Responsable'],
              ticket_rows(), [22, 12, 22, 30, 25, 35, 24, 24, 25, 40, 65, 23, 23, 23, 22, 32, 32, 28, 28])
    sheet('Evolución', ['Inicio del bucket (local)', 'Incidencias', 'Inicio ISO (con offset)'],
          ([bucket.start_at.astimezone(zone).replace(tzinfo=None), bucket.count, bucket.start_at.isoformat()] for bucket in summary.trend), [28, 20, 32])
    for key, title in report_dimensions(summary):
        if key == 'incident_types' and getattr(summary, 'report_kind', None) != 'internal':
            title = 'Tipos'
        if (key == 'sectors' and summary.sector) or (key == 'nodes' and summary.node) or (key == 'responsibles' and summary.responsible):
            continue
        rows = getattr(summary, key)
        if getattr(summary, 'report_kind', None) == 'internal':
            sheet(title, ['ID', 'Nombre', 'Incidencias'], ([row.id, row.label, row.count] for row in rows), [12, 65, 20])
        else:
            sheet(title, ['ID', 'Nombre / código', 'Incidencias', 'Cliente código', 'Cliente nombre'],
                  ([row.id, row.label, row.count, row.customer_code, row.customer_name] for row in rows), [12, 65, 20, 25, 35])
    sheet('Iniciadas y cerradas', ['Bucket ISO', 'Iniciadas', 'Cerradas'], ([row.start_at.isoformat(), row.started, row.closed] for row in summary.activity), [32, 18, 18])
    for key, name in [('hours', 'Horas'), ('weekdays', 'Días')]:
        sheet(name, ['Período', 'Incidencias'], ([row.label, row.count] for row in getattr(summary, key)), [25, 18])
    sheet('Duración por sector', ['Sector', 'Cierres', 'Promedio (segundos)'], ([row.label, row.count, row.average_duration_seconds] for row in summary.sector_durations), [35, 18, 28])
    sheet('Duración por departamento', ['Departamento', 'Cierres', 'Promedio (segundos)'], ([row.label, row.count, row.average_duration_seconds] for row in summary.department_durations), [35, 18, 28])
    if summary.escalations is not None:
        data = summary.escalations
        rows = escalation_metadata(data)
        sheet('Escalamientos', ['Concepto', 'Valor'], rows, [45, 100])
        sheet('Escalamientos por persona', ['ID', 'Destinatario', 'Eventos'], ([row.id, row.label, row.count] for row in data.recipients), [14, 50, 18])
        sheet('Escalamientos por departamento', ['ID', 'Departamento', 'Eventos'], ([row.id, row.label, row.count] for row in data.departments), [14, 50, 18])
        sheet('Evolución escalamientos', ['Bucket ISO', 'Eventos'], ([row.start_at.astimezone(zone).isoformat(), row.count] for row in data.trend), [35, 18])
        sheet('Contexto escalamientos', ['Concepto', 'Valor'], [['Destinatario ID', filters.recipient_id], ['Criterio', 'Inicio operativo del escalamiento en el período; destinatario solo filtra escalamientos.']], [35, 110])
    workbook.save(output)


def export_pdf(session, filters, summary, output):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import KeepTogether, PageBreak, Spacer
    from ticketyn.api.documents import (
        AMBER, BLUE, EMERALD, MUTED, PURPLE, ActivityGraphic, BarsGraphic,
        ModularHeading, TrendGraphic, analytical_block, build_pdf, fonts,
        modern_table, modular_metric_strip, paragraph, report_identification,
        ticket_pdf_table,
    )
    fonts()
    zone = ZoneInfo(filters.timezone)
    width = landscape(A4)[0] - 60
    title = report_title(summary)
    is_internal = getattr(summary, 'report_kind', None) == 'internal'
    accent = PURPLE if is_internal else BLUE
    category = 'Reportería Estadística · NOC Interno' if is_internal else 'Reportería Estadística · Tickets de Clientes'

    period_label = f'{summary.period.from_at.astimezone(zone):%d/%m/%Y %H:%M} — {summary.period.to_exclusive.astimezone(zone):%d/%m/%Y %H:%M} (fin exclusivo)' if summary.period.from_at else 'Todos · Sin restricción temporal'

    context_parts = []
    if not is_internal:
        if summary.sector:
            context_parts.append(f"Sector: {summary.sector['name']}")
        else:
            context_parts.append("Sector: Todos los sectores")
    if summary.node:
        context_parts.append(f"Nodo: {summary.node['name']}")
    if summary.responsible:
        context_parts.append(f"Responsable: {summary.responsible['name']}")
    if is_internal and summary.applied_filters:
        context_parts.extend(summary.applied_filters)
    context_str = ' · '.join(context_parts) if context_parts else 'Todos los registros'

    generated_str = summary.generated_at.astimezone(zone).strftime('%d/%m/%Y %H:%M')

    story = [
        report_identification(
            title,
            period=period_label,
            filters=context_str,
            accent=accent,
            category=category,
            width=width,
        ),
    ]

    primary_kpis = [
        ('INICIADAS', summary.kpis.started, accent),
        ('CERRADAS', summary.kpis.closed, EMERALD),
        ('DURACIÓN PROMEDIO', duration_text(summary.kpis.average_duration_seconds), accent),
        ('DURACIÓN ACUMULADA', duration_text(summary.kpis.total_duration_seconds), accent),
    ]
    story.append(modular_metric_strip(primary_kpis, width, accent=accent))

    if is_internal:
        internal_kpis = [
            ('ABIERTOS (INICIADOS)', summary.open_tickets, AMBER),
            ('CERRADOS (INICIADOS)', summary.closed_tickets, EMERALD),
            ('TICKETS CLIENTES VINCULADOS', summary.related_tickets, PURPLE),
            ('INCIDENCIAS CON VINCULADOS', summary.with_related, PURPLE),
        ]
        story.append(modular_metric_strip(internal_kpis, width, accent=PURPLE))

    story.append(paragraph('Iniciadas, evolución, rankings y detalle: Inicio en el período. Cerradas y duraciones: estado Cerrado y Fin en el período, incluso si comenzaron antes. Cerrados sin Fin no aportan duración. Rankings de Nodo/Responsable: solo asignados.', size=7, color=MUTED))
    story.append(Spacer(1, 6))

    story.append(KeepTogether([
        ModularHeading('Evolución de incidencias', accent=accent, size=10.5),
        TrendGraphic(summary.trend, zone, width, summary.period.granularity, accent=accent),
    ]))
    story.append(Spacer(1, 8))

    story.append(KeepTogether([
        ModularHeading('Iniciadas vs cerradas', accent=accent, size=10.5),
        ActivityGraphic(summary.activity, zone, width, summary.period.granularity, accent=accent),
    ]))
    story.append(Spacer(1, 8))
    number = 4
    for key, name in report_dimensions(summary):
        if (key == 'sectors' and summary.sector) or (key == 'nodes' and summary.node) or (key == 'responsibles' and summary.responsible): continue
        ranking = getattr(summary, key)
        heading = ModularHeading(f'{name} · ranking completo ({len(ranking)})', accent=accent, size=10.5)
        number += 1
        graphic = BarsGraphic(ranking, width, accent=accent) if ranking else None
        if key == 'circuits':
            rows = ([index, row.label, ' — '.join(part for part in [row.customer_code, row.customer_name] if part), row.count] for index, row in enumerate(ranking, 1))
            table = modern_table(['Posición', 'Circuito', 'Cliente', 'Incidencias'], rows, [55, width * .40, width - 55 - width * .40 - 85, 85], light=True, accent=accent)
        else:
            table = modern_table(['Posición', 'Nombre / código', 'Incidencias'],
                ([index, row.label, row.count] for index, row in enumerate(ranking, 1)), [55, width - 140, 85], light=True, accent=accent)
        story.extend(analytical_block(heading, graphic, table, width))
        story.append(Spacer(1, 10))
    story.extend(analytical_block(ModularHeading('Tiempo de resolución por departamento', accent=accent, size=10.5),
        BarsGraphic(summary.department_durations, width, value_field='average_duration_seconds', accent=accent),
        modern_table(['Departamento', 'Cierres', 'Duración promedio'],
            ([row.label, row.count, duration_text(row.average_duration_seconds)] for row in summary.department_durations), [width - 210, 85, 125], light=True, accent=accent), width))
    number += 1
    for key, name in [('hours', 'Distribución por hora del día'), ('weekdays', 'Distribución por día de semana')]:
        values = getattr(summary, key)
        heading, graphic = ModularHeading(name, accent=accent, size=10.5), BarsGraphic(values, width, limit=24, accent=accent)
        if key == 'hours':
            story.append(KeepTogether([heading, graphic]))
        else:
            table = modern_table(['Período', 'Incidencias'], ([row.label, row.count] for row in values), [width - 85, 85], light=True, accent=accent)
            story.extend(analytical_block(heading, graphic, table, width))
        number += 1
    if not summary.sector:
        story.extend(analytical_block(ModularHeading('Duración promedio por sector', accent=accent, size=10.5),
            BarsGraphic(summary.sector_durations, width, value_field='average_duration_seconds', accent=accent),
            modern_table(['Sector', 'Cierres', 'Duración promedio'],
                ([row.label, row.count, duration_text(row.average_duration_seconds)] for row in summary.sector_durations), [width - 210, 85, 125], light=True, accent=accent), width))
        number += 1
    if summary.escalations is not None:
        data = summary.escalations
        story.extend(analytical_block(ModularHeading('Estadísticas de escalamientos', accent=accent, size=10.5), None,
            modern_table(['Indicador', 'Valor'], escalation_metadata(data), [width * .45, width * .55], light=True, accent=accent), width))
        number += 1
        story.append(paragraph(f'Destinatario ID: {filters.recipient_id if filters.recipient_id else "Todos"}. Este filtro no modifica las estadísticas generales de tickets.', size=8, color=MUTED))
        story.extend(analytical_block(ModularHeading('Escalamientos recibidos por persona', accent=accent, size=10.5), BarsGraphic(data.recipients, width, accent=accent),
            modern_table(['Destinatario', 'Eventos'], ([row.label, row.count] for row in data.recipients), [width - 100, 100], light=True, accent=accent), width))
        number += 1
        story.extend(analytical_block(ModularHeading('Escalamientos por departamento', accent=accent, size=10.5), BarsGraphic(data.departments, width, accent=accent), modern_table(['Departamento', 'Eventos'], ([row.label, row.count] for row in data.departments), [width - 100, 100], light=True, accent=accent), width))
        number += 1
        story.append(KeepTogether([ModularHeading('Evolución de escalamientos', accent=accent, size=10.5), TrendGraphic(data.trend, zone, width, data.granularity, accent=accent)]))
        number += 1
    if getattr(summary, 'report_kind', None) == 'internal':
        from ticketyn.api.internal_ticket_exports import export_rows
        table = modern_table(['Correlativo', 'Incidencia', 'Estado', 'Sector', 'Nodo', 'Depto.', 'Responsable', 'Tipo interno', 'Proveedor', 'Inicio', 'Duración', 'Tickets de clientes'], export_rows(session, filters), [width * fraction for fraction in [.10,.13,.06,.07,.07,.07,.09,.09,.09,.085,.065,.08]], reference=True, light=True, accent=PURPLE)
    else:
        table = ticket_pdf_table(report_tickets(session, filters), zone, width, light=True, accent=BLUE)
    story.extend([PageBreak(), ModularHeading('Detalle de tickets — todos los iniciados en el período', accent=accent, size=11, keep=False), table])
    build_pdf(output, title, filters.timezone, story,
              issued_at=generated_str,
              subtitle=category,
              modular=True, landscape_mode=True, reference=title, accent=accent)


def report_title(summary):
    if getattr(summary, 'report_kind', None) == 'internal': return 'Reportería de Tickets internos'
    if summary.node: return "Reporte por Nodo"
    if summary.responsible: return "Reporte por Responsable"
    return "Reporte por Sector" if summary.sector else "Reporte General"


def escalation_metadata(data):
    return [
        ['Tickets únicos escalados', data.escalated_tickets], ['Tickets distintos con eventos en el período', data.event_tickets], ['Eventos totales de escalamiento', data.events],
        ['Escalamientos activos', data.active], ['Escalamientos finalizados', data.finished],
        ['Duración promedio (segundos)', data.average_duration_seconds],
        ['Porcentaje de tickets con escalamiento', data.escalated_percentage],
        ['Tickets del contexto (denominador)', data.total_tickets],
        ['Criterio temporal', 'Inicio operativo del escalamiento. Duración: Fin menos Inicio solo finalizados.'],
        ['Porcentaje', 'Tickets iniciados en el período con escalamiento iniciado en el período / tickets iniciados en el período.'],
    ]


def report_dimensions(summary):
    dimensions = [('sectors', 'Sectores'), ('customers', 'Clientes'), ('circuits', 'Circuitos'), ('incident_types', 'Tipos de incidencia'), ('nodes', 'Nodos'), ('responsibles', 'Responsables'), ('departments', 'Departamentos')]
    if getattr(summary, 'report_kind', None) == 'internal':
        dimensions = [(key, title) for key, title in dimensions if key not in ('customers', 'circuits')]
        dimensions.append(('providers', 'Proveedores'))
    return dimensions
