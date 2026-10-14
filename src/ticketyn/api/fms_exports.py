"""Listados FMS completos, con referencias históricas y sin paginación."""
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP, localcontext
from tempfile import SpooledTemporaryFile
from zoneinfo import ZoneInfo

from fastapi.responses import StreamingResponse
from starlette.background import BackgroundTask
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import Spacer
from sqlalchemy import func, select
from ticketyn.api.documents import build_pdf, fonts, modern_table, paragraph, section
from ticketyn.api.fms_data import operation_query
from ticketyn.api.reports import read_snapshot
from ticketyn.models import Circuit, Customer
from ticketyn.models.fms import Activation, CapacityChange

TITLES = {'activations': 'Activaciones', 'capacity-changes': 'Cambios de capacidad', 'circuit-lifecycle': 'Bajas y Reactivaciones'}
LABELS = {'INCREASE': 'Incremento', 'DECREASE': 'Decremento', 'DEACTIVATE': 'Baja', 'REACTIVATE': 'Reactivación'}


def capacity_text(value, unit):
    if value is None or unit is None:
        return 'Sin registrar'
    with localcontext() as context:
        context.prec = max(28, len(value.as_tuple().digits) + 10)
        mbps = value * (1000 if unit == 'Gbps' else 1)
        output_unit = 'Gbps' if mbps >= 1000 else 'Mbps'
        displayed = (mbps / (1000 if output_unit == 'Gbps' else 1)).quantize(Decimal('.001'), rounding=ROUND_HALF_UP)
        text = format(displayed, 'f').rstrip('0').rstrip('.')
        return f'{text} {output_unit}'


def query_for(model, filters):
    values = filters.model_dump(exclude={'timezone'})
    return operation_query(model, **values)


def friendly_filters(session, filters):
    labels = []
    for field, model, title in [('customer_id', Customer, 'Cliente'), ('circuit_id', Circuit, 'Circuito')]:
        id = getattr(filters, field)
        if id:
            item = session.get(model, id)
            value = ('No disponible' if item is None else f'{item.customer_code} — {item.name}' if model is Customer else item.circuit_code)
            labels.append(f'{title}: {value}')
    operation_type = getattr(filters, 'operation_type', None)
    if operation_type:
        labels.append('Tipo: ' + LABELS[operation_type])
    zone = ZoneInfo(filters.timezone)
    for field, title in [('from_at', 'Desde (incluido)'), ('to_at', 'Hasta (excluido)')]:
        value = getattr(filters, field)
        if value:
            labels.append(f'{title}: {value.astimezone(zone).isoformat()}')
    return labels or ['Todos los registros']


def headers(model, exact=False):
    common = ['Fecha y hora efectiva', 'Cliente', 'Código del circuito', 'Descripción del circuito', 'Sector', 'Nodo', 'Responsable']
    if model is Activation:
        return ['Referencia de activación', *common]
    if model is CapacityChange:
        return [*common, 'Tipo de operación', 'Capacidad anterior', 'Capacidad nueva', 'Descripción', *(['Valor anterior registrado (exacto)', 'Unidad anterior registrada', 'Valor nuevo registrado (exacto)', 'Unidad nueva registrada'] if exact else [])]
    return [*common, 'Tipo de operación', 'Motivo de baja', 'Observaciones']


def rows(session, model, filters, exact=False, pdf=False):
    zone = ZoneInfo(filters.timezone)
    # Todas las relaciones escalares se cargan en el mismo SELECT; no bitácoras.
    for row in session.scalars(query_for(model, filters).execution_options(yield_per=200)):
        circuit = row.circuit
        effective = row.effective_at.astimezone(zone)
        values = [effective.strftime('%d/%m/%Y\n%H:%M:%S') if pdf else effective.isoformat(), circuit.customer.name, circuit.circuit_code,
                  circuit.description or 'Sin descripción', row.sector.name if row.sector else None,
                  row.node.name if row.node else None, row.responsible.name]
        if model is Activation:
            values = [row.reference, *values]
        elif model is CapacityChange:
            values += [LABELS[row.operation_type], capacity_text(row.previous_capacity, row.previous_unit), capacity_text(row.new_capacity, row.new_unit), row.description]
            if exact:
                values += [str(row.previous_capacity) if row.previous_capacity is not None else None, row.previous_unit, str(row.new_capacity), row.new_unit]
        else:
            values += [LABELS[row.operation_type], row.reason.name if row.operation_type == 'DEACTIVATE' and row.reason else None, row.observations]
        yield values


def export_pdf(session, model, kind, filters, output):
    fonts()
    title = TITLES[kind]
    zone = ZoneInfo(filters.timezone)
    generated = datetime.now(timezone.utc).astimezone(zone)
    total = session.scalar(select(func.count()).select_from(query_for(model, filters).order_by(None).subquery()))
    columns = headers(model)
    # Proporcionar más espacio a descripciones/observaciones y respetar A4 horizontal.
    weights = [2 if column in ('Descripción del circuito', 'Descripción', 'Observaciones') else 1.3 if column in ('Cliente', 'Responsable', 'Fecha y hora efectiva') else 1.2 if column in ('Tipo de operación', 'Código del circuito') else 1 for column in columns]
    width = landscape(A4)[0] - 60
    pdf_columns = [column.replace('Código del circuito', 'Código del\ncircuito').replace('Tipo de operación', 'Tipo de\noperación').replace('Capacidad anterior', 'Capacidad\nanterior').replace('Capacidad nueva', 'Capacidad\nnueva') for column in columns]
    table = modern_table(pdf_columns, ([value if value is not None else 'Sin registrar' for value in record] for record in rows(session, model, filters, pdf=True)), [width * weight / sum(weights) for weight in weights], reference=model is Activation, compact=True)
    story = [paragraph(title, size=21, bold=True), paragraph(f'Generado: {generated.astimezone(zone):%d/%m/%Y %H:%M} · Total encontrado: {total}', size=9), paragraph(' · '.join(friendly_filters(session, filters)), size=8), Spacer(1, 10), section(title, 1, keep=False), table]
    build_pdf(output, title, filters.timezone, story, issued_at=generated)


def export_xlsx(session, model, kind, filters, output):
    book = Workbook(write_only=True)
    sheet = book.create_sheet(TITLES[kind]); sheet.freeze_panes = 'A2'
    columns = headers(model, exact=True)
    for i, title in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(i)].width = 48 if title in ('Descripción', 'Descripción del circuito', 'Observaciones') else 30
    def append(target, values, heading=False):
        cells = []
        for value in values:
            cell = WriteOnlyCell(target, value=value)
            if isinstance(value, str): cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if heading:
                cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='102744')
            cells.append(cell)
        target.append(cells)
    append(sheet, columns, True)
    count = 1
    for record in rows(session, model, filters, exact=True):
        append(sheet, record); count += 1
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{count}'
    metadata = book.create_sheet('Resumen')
    append(metadata, ['Dato', 'Valor'], True)
    append(metadata, ['Módulo', TITLES[kind]])
    append(metadata, ['Generado', datetime.now(timezone.utc).astimezone(ZoneInfo(filters.timezone)).isoformat()])
    append(metadata, ['Zona horaria', filters.timezone]); append(metadata, ['Total exportado', count - 1])
    for value in friendly_filters(session, filters): append(metadata, ['Filtro', value])
    book.save(output)


def stream_export(session, model, kind, format, filters):
    read_snapshot(session)
    output = SpooledTemporaryFile(max_size=2 * 1024 * 1024)
    try:
        (export_pdf if format == 'pdf' else export_xlsx)(session, model, kind, filters, output)
        output.seek(0)
    except Exception:
        output.close()
        raise
    def chunks():
        try:
            while chunk := output.read(65536): yield chunk
        finally:
            output.close()
    media = 'application/pdf' if format == 'pdf' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return StreamingResponse(chunks(), media_type=media, headers={'Content-Disposition': f'attachment; filename="ticketyn-{kind}.{format}"', 'Cache-Control': 'no-store'}, background=BackgroundTask(output.close))
