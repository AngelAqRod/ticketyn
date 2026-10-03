"""Document composition only; report data and XLSX remain unchanged."""
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from reportlab.platypus import KeepTogether, Spacer

from ticketyn.api.documents import (
    analytical_block, BarsGraphic, build_pdf, fonts, modern_table, section,
)


@pytest.mark.parametrize('count, grouped', [(1, True), (7, True), (150, False)])
def test_measured_ranking_blocks_can_paginate(count, grouped):
    fonts()
    rows = [SimpleNamespace(label=f'Categoría {i:03}', count=i) for i in range(count)]
    width = 780
    table = modern_table(['Nombre', 'Incidencias'], ([row.label, row.count] for row in rows), [650, 130])
    block = analytical_block(section('Ranking de prueba'), BarsGraphic(rows, width), table, width)
    assert (len(block) == 1 and isinstance(block[0], KeepTogether)) == grouped
    assert table.repeatRows == 1
    output = BytesIO()
    build_pdf(output, 'Prueba de composición', 'UTC', [Spacer(1, 320), *block])
    assert output.getvalue().startswith(b'%PDF')


def test_pdf_distributions_and_small_rankings(api_client, ticket_payload, monkeypatch):
    from reportlab.pdfgen.textobject import PDFTextObject
    node = api_client.post('/api/nodes', json={'name': 'Nodo composición'}).json()
    responsible = api_client.post('/api/responsibles', json={'name': 'Responsable composición'}).json()
    api_client.patch(f"/api/circuits/{ticket_payload['circuit_id']}", json={'node_id': node['id']})
    api_client.post('/api/tickets', json={**ticket_payload, 'responsible_id': responsible['id']})
    params = {'from': '2020-01-01T00:00:00Z', 'to': '2020-01-02T00:00:00Z'}
    summary = api_client.get('/api/reports/summary', params=params).json()
    assert len(summary['hours']) == 24
    assert len(summary['weekdays']) == 7
    tables = []
    # export_pdf imports the shared table helper locally.
    from ticketyn.api import documents
    original_table = documents.modern_table
    def capture_table(headers, rows, widths, **kwargs):
        rows = list(rows)
        tables.append((headers, rows))
        return original_table(headers, rows, widths, **kwargs)
    monkeypatch.setattr(documents, 'modern_table', capture_table)
    texts = []
    footers = []
    from reportlab.pdfgen.canvas import Canvas
    original_right = Canvas.drawRightString
    def capture_right(self, x, y, text, *args, **kwargs):
        footers.append(text)
        return original_right(self, x, y, text, *args, **kwargs)
    monkeypatch.setattr(Canvas, "drawRightString", capture_right)
    original_text = PDFTextObject._textOut
    def capture_text(self, value, *args, **kwargs):
        texts.append((self._canvas._pageNumber, value))
        return original_text(self, value, *args, **kwargs)
    monkeypatch.setattr(PDFTextObject, '_textOut', capture_text)
    response = api_client.get('/api/reports/export/pdf', params=params)
    assert response.status_code == 200
    assert not any(headers == ['Período', 'Incidencias'] and len(rows) == 24 for headers, rows in tables)
    assert any(headers == ['Período', 'Incidencias'] and len(rows) == 7 for headers, rows in tables)
    for label in ['Nodo composición', 'Responsable composición']:
        pages = {page for page, value in texts if label in value}
        assert len(pages) == 1  # Chart and complete small table travel together.
    assert any('Detalle de tickets' in value for _, value in texts)
    assert any(value.startswith('Página ') for value in footers)
    assert any('Reporte General' in value for value in footers)
    book = load_workbook(BytesIO(api_client.get('/api/reports/export/xlsx', params=params).content))
    assert book['Horas'].max_row == 25
    assert book['Días'].max_row == 8
