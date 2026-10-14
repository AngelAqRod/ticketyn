"""Recovery of missing history, exclusively on disposable PostgreSQL."""
from uuid import uuid4
from pathlib import Path
import runpy
import pytest
from sqlalchemy import create_engine, text
from alembic.migration import MigrationContext
from alembic.operations import Operations
from test_e2e_update_migration import cli
from test_ticket_exports import pdf_strings


@pytest.fixture
def legacy_nodes(postgres_engine):
    name = 'recover_nodes_' + uuid4().hex
    with postgres_engine.connect().execution_options(isolation_level='AUTOCOMMIT') as admin:
        admin.execute(text('CREATE DATABASE ' + name))
    engine = create_engine(postgres_engine.url.set(database=name))
    try:
        cli(engine.url, 'upgrade', '0017_historical_nodes')
        with engine.begin() as db:
            db.execute(text("INSERT INTO customers(id,customer_code,name) VALUES(1,'ADOPT-TEST-001','Cliente')"))
            db.execute(text("INSERT INTO nodes(id,name) VALUES(1,'UFINET'),(2,'Histórico conservado')"))
            for table in ('sectors','departments','incident_types','responsibles'):
                db.execute(text(f"INSERT INTO {table}(id,name) VALUES(1,'Ejemplo')"))
            db.execute(text("INSERT INTO circuits(id,customer_id,circuit_code,description,node_id,sector_id) VALUES(1,1,'ADOPT-TEST-001.01','Con Nodo',1,1),(2,1,'SIN-NODO','Sin Nodo',NULL,1)"))
            for tid,cid,nid in [(1,1,None),(2,1,2),(3,2,None),(4,1,None)]:
                db.execute(text("INSERT INTO tickets(id,ticket_number,reference,title,description,customer_id,circuit_id,node_id,sector_id,department_id,incident_type_id,start_at,status) VALUES(:id,:id,:ref,'Histórico','Datos',1,:cid,:nid,1,1,1,'2020-01-01T00:00:00Z','OPEN')"), {'id':tid,'cid':cid,'nid':nid,'ref':str(tid)})
            db.execute(text("INSERT INTO capacity_changes(circuit_id,responsible_id,effective_at,operation_type,new_capacity,new_unit) VALUES(1,1,'2020-01-01T00:00:00Z','INCREASE',500,'Mbps')"))
            db.execute(text("INSERT INTO circuit_lifecycle_operations(circuit_id,responsible_id,effective_at,operation_type) VALUES(1,1,'2020-01-01T00:00:00Z','REACTIVATE')"))
            db.execute(text("INSERT INTO activations(circuit_id,responsible_id,effective_at,activation_number,reference) VALUES(1,1,'2020-01-01T00:00:00Z',1,'ACT-1')"))
        yield engine
    finally:
        engine.dispose()
        with postgres_engine.connect().execution_options(isolation_level='AUTOCOMMIT') as admin:
            admin.execute(text('DROP DATABASE ' + name))


def apply_again(engine):
    migration = runpy.run_path(str(Path(__file__).resolve().parents[1] / 'alembic/versions/0018_recover_ticket_nodes.py'))
    with engine.begin() as db:
        with Operations.context(MigrationContext.configure(db)):
            migration['upgrade']()


def test_recovery_preserves_everything_else_and_is_idempotent(legacy_nodes):
    engine = legacy_nodes
    with engine.connect() as db:
        before = db.execute(text('SELECT to_jsonb(t) FROM tickets t ORDER BY id')).scalars().all()
        circuits = db.execute(text('SELECT to_jsonb(c) FROM circuits c ORDER BY id')).scalars().all()
        fms = {table: db.execute(text(f'SELECT to_jsonb(a) FROM {table} a ORDER BY id')).scalars().all() for table in ('activations','capacity_changes','circuit_lifecycle_operations')}
        eligible = db.scalar(text('SELECT count(*) FROM tickets t JOIN circuits c ON c.id=t.circuit_id JOIN nodes n ON n.id=c.node_id WHERE t.node_id IS NULL'))
        assert eligible == 2
    cli(engine.url, 'upgrade', '0018_recover_ticket_nodes')
    apply_again(engine)
    with engine.connect() as db:
        after = db.execute(text('SELECT to_jsonb(t) FROM tickets t ORDER BY id')).scalars().all()
        assert [x['node_id'] for x in after] == [1,2,None,1]
        assert [{k:v for k,v in x.items() if k != 'node_id'} for x in after] == [{k:v for k,v in x.items() if k != 'node_id'} for x in before]
        assert db.execute(text('SELECT to_jsonb(c) FROM circuits c ORDER BY id')).scalars().all() == circuits
        for table, original in fms.items():
            assert db.execute(text(f'SELECT to_jsonb(a) FROM {table} a ORDER BY id')).scalars().all() == original
        assert db.scalar(text('SELECT version_num FROM alembic_version')) == '0018_recover_ticket_nodes'
    cli(engine.url, 'downgrade', '0017_historical_nodes')
    with engine.connect() as db:
        assert db.execute(text('SELECT node_id FROM tickets ORDER BY id')).scalars().all() == [1,2,None,1]


@pytest.mark.parametrize('corruption', ['orphan_circuit', 'invalid_node', 'invalid_ticket_node', 'no_circuit'])
def test_inconsistent_relations_fail_before_any_assignment(legacy_nodes, corruption):
    engine = legacy_nodes
    with engine.begin() as db:
        if corruption in ('orphan_circuit','no_circuit'):
            db.execute(text('ALTER TABLE tickets DROP CONSTRAINT fk_tickets_circuit_id_circuits'))
            db.execute(text('ALTER TABLE tickets ALTER COLUMN circuit_id DROP NOT NULL'))
            db.execute(text('UPDATE tickets SET circuit_id=:value WHERE id=4'), {'value':999 if corruption == 'orphan_circuit' else None})
        elif corruption == 'invalid_ticket_node':
            db.execute(text('ALTER TABLE tickets DROP CONSTRAINT fk_tickets_node_id_nodes'))
            db.execute(text('UPDATE tickets SET node_id=999 WHERE id=2'))
        else:
            db.execute(text('ALTER TABLE circuits DROP CONSTRAINT fk_circuits_node_id_nodes'))
            db.execute(text('UPDATE circuits SET node_id=999 WHERE id=2'))
    if corruption == 'no_circuit':
        apply_again(engine)
        with engine.connect() as db:
            assert db.scalar(text('SELECT node_id FROM tickets WHERE id=4')) is None
    else:
        with pytest.raises(RuntimeError, match='inconsistentes'):
            apply_again(engine)
        with engine.connect() as db:
            assert db.scalar(text('SELECT node_id FROM tickets WHERE id=1')) is None
            assert db.scalar(text('SELECT version_num FROM alembic_version')) == '0017_historical_nodes'


def test_existing_creation_and_historical_consumers(api_client, ticket_payload, db_session, pdf_strings):
    from ticketyn.models import Circuit, Ticket
    from test_circuit_nodes import node
    body = {key:value for key,value in ticket_payload.items() if key != 'node_id'}
    created = api_client.post('/api/tickets', json=body)
    assert created.status_code == 201
    ticket = created.json()
    old = ticket_payload['node_id']
    assert ticket['node_id'] == old
    current = node(api_client, 'NODO ACTUAL POSTERIOR')
    assert api_client.patch(f"/api/circuits/{body['circuit_id']}", json={'node_id': current}).status_code == 200
    assert api_client.get(f"/api/tickets/{ticket['id']}").json()['node_id'] == old
    assert api_client.get('/api/tickets', params={'node_id': old}).json()[0]['node_id'] == old
    assert api_client.get('/api/tickets', params={'node_id': current}).json() == []
    assert api_client.get('/api/reports/summary', params={'node_id': old}).json()['kpis']['started'] == 1
    for path in ['/api/tickets/export/pdf', f"/api/tickets/{ticket['id']}/reports/internal/pdf", '/api/reports/export/pdf']:
        pdf_strings.clear()
        response = api_client.get(path, params={'node_id': old} if '/tickets/' + str(ticket['id']) not in path else {})
        assert response.status_code == 200, response.text
        text_content = ' '.join(pdf_strings)
        assert 'Nodo de prueba' in text_content
        assert 'NODO ACTUAL POSTERIOR' not in text_content
    # A legacy NULL remains explicit in the API before a data repair, never a visual fallback.
    historical = db_session.get(Ticket, ticket['id']); historical.node_id = None
    circuit = db_session.get(Circuit, body['circuit_id']); circuit.node_id = None
    db_session.commit()
    assert api_client.get(f"/api/tickets/{ticket['id']}").json()['node'] is None
    result = api_client.post('/api/tickets', json={**body, 'node_id': old})
    assert result.status_code == 201
    assert result.json()['node_id'] == old
    assert api_client.get(f"/api/circuits/{body['circuit_id']}").json()['node_id'] == old


def test_repaired_rows_are_used_by_real_api_and_exports(legacy_nodes, api_client, pdf_strings):
    from sqlalchemy.orm import Session
    from ticketyn.db.session import get_session
    from io import BytesIO
    from openpyxl import load_workbook
    engine = legacy_nodes
    def session():
        with Session(engine) as db:
            yield db
    api_client.app.dependency_overrides[get_session] = session
    # Runtime ORM uses the current schema; reproduce legacy NULLs on it.
    cli(engine.url, 'upgrade', 'head')
    with engine.begin() as db:
        db.execute(text('UPDATE tickets SET node_id=NULL WHERE id IN (1,4)'))
    assert api_client.get('/api/tickets/1').json()['node'] is None
    assert api_client.get('/api/circuits/1').json()['node']['name'] == 'UFINET'
    apply_again(engine)
    assert api_client.get('/api/tickets/1').json()['node']['name'] == 'UFINET'
    assert {x['id'] for x in api_client.get('/api/tickets', params={'node_id': 1}).json()} == {1,4}
    assert api_client.get('/api/tickets/2').json()['node_id'] == 2
    assert api_client.get('/api/reports/summary', params={'node_id': 1}).json()['kpis']['started'] == 2
    book = load_workbook(BytesIO(api_client.get('/api/tickets/export/xlsx', params={'node_id': 1}).content))
    sheet = book['Tickets']
    column = [x.value for x in sheet[1]].index('Nodo') + 1
    assert sheet.max_row == 3
    assert {sheet.cell(row,column).value for row in range(2,4)} == {'UFINET'}
    assert api_client.get('/api/tickets/1/reports/internal/pdf').status_code == 200
    assert 'UFINET' in ' '.join(pdf_strings)
    assert api_client.patch('/api/circuits/1', json={'node_id': 2}).status_code == 200
    assert api_client.get('/api/tickets/1').json()['node_id'] == 1
    assert api_client.get('/api/activations/1').json()['node'] is None
