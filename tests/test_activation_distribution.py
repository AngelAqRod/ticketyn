"""Pruebas de equipo y puerto de distribución en activaciones FMS (Ticketyn v0.8.0)."""
from datetime import datetime, timezone
from io import BytesIO
from uuid import uuid4
import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, inspect, text
from sqlalchemy.exc import IntegrityError

from test_e2e_update_migration import cli
from test_fms import create, operation_payload
from test_ticket_exports import pdf_strings
from ticketyn.api.fms_data import distribution_labels
from ticketyn.models.fms import Activation


@pytest.fixture
def distribution_setup(api_client, ticket_payload, numbered_port):
    node_id = ticket_payload['node_id']
    device = api_client.post('/api/distribution-devices', json={
        'node_id': node_id,
        'name': 'OLT-HUAWEI-01',
        'description': 'Cabecera principal de distribución'
    }).json()
    port = numbered_port(device['id'], prefix='gpon-', numbering_format='X/Y/Z', positions=(0, 1), description='Puerto troncal pon 1')
    return {'node_id': node_id, 'device': device, 'port': port}


def test_1_activation_with_device_and_port(api_client, operation_payload, distribution_setup):
    """1. Activación con equipo y puerto asignados."""
    cid = operation_payload['circuit_id']
    dev = distribution_setup['device']
    port = distribution_setup['port']
    patch_res = api_client.patch(f'/api/circuits/{cid}', json={
        'distribution_device_id': dev['id'],
        'distribution_port_id': port['id']
    })
    assert patch_res.status_code == 200

    act = create(api_client, '/api/activations', operation_payload)
    assert act['distribution_recorded'] is True
    assert act['distribution_device_name'] == 'OLT-HUAWEI-01'
    assert act['distribution_port_name'] == 'gpon-0/1/1'

    fetched = api_client.get(f"/api/activations/{act['id']}").json()
    assert fetched['distribution_recorded'] is True
    assert fetched['distribution_device_name'] == 'OLT-HUAWEI-01'
    assert fetched['distribution_port_name'] == 'gpon-0/1/1'


def test_2_activation_with_device_without_port(api_client, operation_payload, distribution_setup):
    """2. Activación con equipo de distribución pero sin puerto asignado."""
    cid = operation_payload['circuit_id']
    dev = distribution_setup['device']
    patch_res = api_client.patch(f'/api/circuits/{cid}', json={
        'distribution_device_id': dev['id'],
        'distribution_port_id': None
    })
    assert patch_res.status_code == 200

    act = create(api_client, '/api/activations', operation_payload)
    assert act['distribution_recorded'] is True
    assert act['distribution_device_name'] == 'OLT-HUAWEI-01'
    assert act['distribution_port_name'] is None

    fetched = api_client.get(f"/api/activations/{act['id']}").json()
    assert fetched['distribution_device_name'] == 'OLT-HUAWEI-01'
    assert fetched['distribution_port_name'] is None


def test_3_activation_without_assignment(api_client, operation_payload):
    """3. Activación sin asignación de infraestructura de distribución («No asignado»)."""
    cid = operation_payload['circuit_id']
    patch_res = api_client.patch(f'/api/circuits/{cid}', json={
        'distribution_device_id': None,
        'distribution_port_id': None
    })
    assert patch_res.status_code == 200

    act = create(api_client, '/api/activations', operation_payload)
    assert act['distribution_recorded'] is True
    assert act['distribution_device_name'] is None
    assert act['distribution_port_name'] is None

    fetched = api_client.get(f"/api/activations/{act['id']}").json()
    assert fetched['distribution_recorded'] is True
    assert fetched['distribution_device_name'] is None
    assert fetched['distribution_port_name'] is None


def test_4_subsequent_circuit_distribution_change_does_not_rewrite_history(api_client, operation_payload, distribution_setup):
    """4. Modificación posterior del equipo/puerto del circuito no reescribe la activación."""
    cid = operation_payload['circuit_id']
    dev = distribution_setup['device']
    port = distribution_setup['port']
    api_client.patch(f'/api/circuits/{cid}', json={
        'distribution_device_id': dev['id'],
        'distribution_port_id': port['id']
    })

    act = create(api_client, '/api/activations', operation_payload)
    assert act['distribution_device_name'] == 'OLT-HUAWEI-01'
    assert act['distribution_port_name'] == 'gpon-0/1/1'

    # Cambiar circuito a otro equipo o retirarlo
    dev2 = api_client.post('/api/distribution-devices', json={
        'node_id': distribution_setup['node_id'],
        'name': 'SWITCH-CISCO-02'
    }).json()
    api_client.patch(f'/api/circuits/{cid}', json={
        'distribution_device_id': dev2['id'],
        'distribution_port_id': None
    })

    # La activación histórica previa permanece intacta
    persisted = api_client.get(f"/api/activations/{act['id']}").json()
    assert persisted['distribution_device_name'] == 'OLT-HUAWEI-01'
    assert persisted['distribution_port_name'] == 'gpon-0/1/1'


def test_5_historical_activations_prior_to_feature_show_not_registered(api_client, operation_payload, distribution_setup, db_session):
    """5. Activaciones históricas previas representan la distribución desconocida como «No registrado»."""
    cid = operation_payload['circuit_id']
    dev = distribution_setup['device']
    api_client.patch(f'/api/circuits/{cid}', json={'distribution_device_id': dev['id']})

    # Simular una activación legacy creada antes de esta característica (distribution_recorded = False)
    db_session.execute(text("""
        INSERT INTO activations (
            circuit_id, responsible_id, effective_at, activation_number, reference,
            distribution_recorded, distribution_device_name, distribution_port_name
        ) VALUES (
            :cid, :rid, '2020-01-01T00:00:00Z', 99999, 'ACT-HISTORICA',
            false, NULL, NULL
        )
    """), {'cid': cid, 'rid': operation_payload['responsible_id']})
    db_session.commit()

    legacy = api_client.get('/api/activations', params={'circuit_id': cid}).json()
    act = next(a for a in legacy if a['reference'] == 'ACT-HISTORICA')
    assert act['distribution_recorded'] is False
    assert act['distribution_device_name'] is None
    assert act['distribution_port_name'] is None

    # Helper de backend debe devolver explícitamente ('No registrado', 'No registrado')
    db_act = db_session.get(Activation, act['id'])
    assert distribution_labels(db_act) == ('No registrado', 'No registrado')


def test_6_pdf_and_xlsx_exports_contain_distribution_fields(api_client, operation_payload, distribution_setup, pdf_strings):
    """6. Exportaciones PDF y XLSX reflejan Equipo de distribución y Puerto."""
    cid = operation_payload['circuit_id']
    dev = distribution_setup['device']
    port = distribution_setup['port']
    api_client.patch(f'/api/circuits/{cid}', json={
        'distribution_device_id': dev['id'],
        'distribution_port_id': port['id']
    })
    act = create(api_client, '/api/activations', operation_payload)

    # 6a. PDF individual
    pdf_strings.clear()
    rep = api_client.get(f"/api/activations/{act['id']}/reports/pdf")
    assert rep.status_code == 200
    content = ' '.join(pdf_strings)
    assert 'Equipo de distribución' in content
    assert 'Puerto' in content
    assert 'OLT-HUAWEI-01' in content
    assert 'gpon-0/1/1' in content

    # 6b. Exportación tabular XLSX
    xlsx_res = api_client.get('/api/activations/export/xlsx', params={'circuit_id': cid})
    assert xlsx_res.status_code == 200
    book = load_workbook(BytesIO(xlsx_res.content))
    sheet = book.worksheets[0]
    headers = [cell.value for cell in sheet[1]]
    assert 'Equipo de distribución' in headers
    assert 'Puerto' in headers
    data = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    target_row = next(r for r in data if r['Referencia de activación'] == act['reference'])
    assert target_row['Equipo de distribución'] == 'OLT-HUAWEI-01'
    assert target_row['Puerto'] == 'gpon-0/1/1'
    book.close()

    # 6c. Exportación tabular PDF
    pdf_strings.clear()
    tab_pdf = api_client.get('/api/activations/export/pdf', params={'circuit_id': cid})
    assert tab_pdf.status_code == 200
    tab_content = ' '.join(pdf_strings)
    assert 'OLT-HUAWEI-01' in tab_content
    assert 'gpon-0/1/1' in tab_content


def test_7_data_integrity_and_database_constraints(db_session, operation_payload):
    """7. Restricciones de integridad en base de datos para snapshot de distribución."""
    cid = operation_payload['circuit_id']
    rid = operation_payload['responsible_id']

    # Puerto sin equipo de distribución debe fallar por CheckConstraint distribution_device_required
    with pytest.raises(IntegrityError):
        with db_session.begin_nested():
            db_session.execute(text("""
                INSERT INTO activations (
                    circuit_id, responsible_id, effective_at, activation_number, reference,
                    distribution_recorded, distribution_device_name, distribution_port_name
                ) VALUES (
                    :cid, :rid, '2020-01-01T00:00:00Z', 90001, 'ACT-INVALID-1',
                    true, NULL, 'puerto-huerfano'
                )
            """), {'cid': cid, 'rid': rid})

    # distribution_recorded = false pero con nombres poblados debe fallar por distribution_recorded_valid
    with pytest.raises(IntegrityError):
        with db_session.begin_nested():
            db_session.execute(text("""
                INSERT INTO activations (
                    circuit_id, responsible_id, effective_at, activation_number, reference,
                    distribution_recorded, distribution_device_name, distribution_port_name
                ) VALUES (
                    :cid, :rid, '2020-01-01T00:00:00Z', 90002, 'ACT-INVALID-2',
                    false, 'Equipo fantasma', NULL
                )
            """), {'cid': cid, 'rid': rid})


def test_0034_upgrade_downgrade_preserves_existing_data(postgres_engine):
    """Verificación de upgrade y downgrade de migración 0034 en clúster PostgreSQL efímero."""
    name = 'act_dist_mig_' + uuid4().hex
    with postgres_engine.connect().execution_options(isolation_level='AUTOCOMMIT') as connection:
        connection.execute(text('CREATE DATABASE ' + name))
    engine = create_engine(postgres_engine.url.set(database=name))
    try:
        # 1. Migrar hasta 0033
        cli(engine.url, 'upgrade', '0033_distribution_devices_ports')
        with engine.begin() as conn:
            conn.execute(text("INSERT INTO customers(id, name, customer_code) VALUES (1, 'Cliente', 'CLI')"))
            conn.execute(text("INSERT INTO responsibles(id, name) VALUES (1, 'Resp')"))
            conn.execute(text("INSERT INTO circuits(id, customer_id, circuit_code, description) VALUES (1, 1, 'C-1', 'Desc')"))
            conn.execute(text("""
                INSERT INTO activations (circuit_id, responsible_id, effective_at, activation_number, reference)
                VALUES (1, 1, '2020-01-01T00:00:00Z', 1, 'ACT-001')
            """))

        # 2. Verificar específicamente 0034, independientemente de revisiones posteriores.
        cli(engine.url, 'upgrade', '0034_activation_distribution')
        with engine.connect() as conn:
            assert conn.scalar(text('SELECT version_num FROM alembic_version')) == '0034_activation_distribution'
            cols = {c['name']: c for c in inspect(conn).get_columns('activations')}
            assert 'distribution_recorded' in cols
            assert 'distribution_device_name' in cols
            assert 'distribution_port_name' in cols

            # Registro anterior conserva valores nulos y distribution_recorded = False
            row = conn.execute(text("SELECT distribution_recorded, distribution_device_name, distribution_port_name FROM activations WHERE activation_number = 1")).mappings().one()
            assert row['distribution_recorded'] is False
            assert row['distribution_device_name'] is None
            assert row['distribution_port_name'] is None

        # 3. Intentar downgrade cuando hay registros con distribution_recorded = true debe fallar
        with engine.begin() as conn:
            conn.execute(text("""
                INSERT INTO activations (circuit_id, responsible_id, effective_at, activation_number, reference, distribution_recorded, distribution_device_name)
                VALUES (1, 1, '2020-01-02T00:00:00Z', 2, 'ACT-002', true, 'SW-01')
            """))

        with pytest.raises(Exception) as excinfo:
            cli(engine.url, 'downgrade', '0033_distribution_devices_ports')
        assert 'Downgrade inseguro' in str(excinfo.value) or 'Downgrade inseguro' in repr(excinfo.value)

        # 4. Eliminar el registro protegido y ejecutar downgrade exitoso
        with engine.begin() as conn:
            conn.execute(text("DELETE FROM activations WHERE activation_number = 2"))

        cli(engine.url, 'downgrade', '0033_distribution_devices_ports')
        with engine.connect() as conn:
            assert conn.scalar(text('SELECT version_num FROM alembic_version')) == '0033_distribution_devices_ports'
            cols = {c['name'] for c in inspect(conn).get_columns('activations')}
            assert 'distribution_recorded' not in cols
            assert 'distribution_device_name' not in cols
            assert 'distribution_port_name' not in cols
    finally:
        engine.dispose()
        with postgres_engine.connect().execution_options(isolation_level='AUTOCOMMIT') as connection:
            connection.execute(text('DROP DATABASE ' + name))
