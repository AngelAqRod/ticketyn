from io import BytesIO
import pytest
from openpyxl import load_workbook
from test_internal_tickets import payload
from test_ticket_exports import pdf_strings

ROOT='/api/reports/internal-tickets'

def summary(client, **filters):
    response=client.get(ROOT+'/summary',params=filters)
    assert response.status_code==200,response.text
    return response.json()


def test_empty_and_exports(api_client,pdf_strings):
    result=summary(api_client)
    assert result['kpis']['started']==0 and result['related_tickets']==0
    assert result['kpis']['average_duration_seconds'] is None
    for format in ('pdf','xlsx'):
        response=api_client.get(ROOT+'/export/'+format)
        assert response.status_code==200,response.text
        assert 'tickets-internos' in response.headers['content-disposition']
        if format=='pdf':
            assert response.content.startswith(b'%PDF-')
            assert 'Reportería de Tickets internos' in ' '.join(pdf_strings)
        else:
            book=load_workbook(BytesIO(response.content))
            assert 'Tickets internos' in book.sheetnames and 'Proveedores' in book.sheetnames
            assert 'Clientes' not in book.sheetnames and 'Circuitos' not in book.sheetnames


def test_counts_provider_history_and_relations(api_client,ticket_payload,pdf_strings):
    p=payload(api_client)
    provider=api_client.post('/api/providers',json={'name':'Transporte histórico'}).json()
    a=api_client.post('/api/internal-tickets',json={**p,'provider_id':provider['id'],'status':'CLOSED','end_at':'2026-10-15T10:00:00-06:00'}).json()
    b=api_client.post('/api/internal-tickets',json=p).json()
    for _ in range(5):
        assert api_client.post('/api/tickets',json={**ticket_payload,'internal_ticket_id':a['id']}).status_code==201
    api_client.patch(f"/api/providers/{provider['id']}",json={'active':False})
    result=summary(api_client)
    assert result['kpis']=={'started':2,'closed':1,'average_duration_seconds':7200.0,'total_duration_seconds':7200.0}
    assert (result['open_tickets'],result['closed_tickets'])==(1,1)
    assert (result['related_tickets'],result['with_related'],result['without_related'])==(5,1,1)
    assert {r['label']:r['count'] for r in result['providers']}=={'Transporte histórico':1,'Ninguno':1}
    assert result['incident_types'][0]['label']=='Energía'
    assert result['customers']==[] and result['circuits']==[]
    assert summary(api_client,provider_id=provider['id'])['kpis']['started']==1
    assert summary(api_client,without_provider=True)['related_tickets']==0
    assert summary(api_client,status='OPEN')['kpis']['closed']==0
    for key in ('sector_id','department_id','incident_type_id'):
        assert summary(api_client,**{key:p[key]})['kpis']['started']==2
        if key != 'sector_id':
            assert summary(api_client,**{key:99999})['kpis']['started']==0
    pdf=api_client.get(ROOT+'/export/pdf',params={'provider_id':provider['id']})
    assert pdf.status_code==200,pdf.text
    text=' '.join(pdf_strings)
    assert 'Transporte histórico' in text and 'Tickets de clientes vinculados' in text
    assert a['reference'] in text and b['reference'] not in text
    excel=api_client.get(ROOT+'/export/xlsx',params={'provider_id':provider['id']})
    assert excel.status_code==200,excel.text
    book=load_workbook(BytesIO(excel.content))
    rows=list(book['Tickets internos'].values)
    assert len(rows)==2 and rows[1][-1]==5
    normal=api_client.get('/api/reports/summary').json()
    assert normal['kpis']['started']==5


def test_period_uses_start_and_end_separately(api_client):
    p=payload(api_client)
    api_client.post('/api/internal-tickets',json={**p,'start_at':'2026-10-14T08:00:00-06:00','end_at':'2026-10-15T10:00:00-06:00','status':'CLOSED'})
    result=summary(api_client,**{'from':'2026-10-15T00:00:00-06:00','to':'2026-10-16T00:00:00-06:00','timezone':'America/Guatemala'})
    assert result['kpis']['started']==0 and result['kpis']['closed']==1
    # Cerrado en el período, pero iniciado antes: métricas deliberadamente distintas.
    assert result['closed_tickets']==0
    assert result['kpis']['average_duration_seconds']==26*3600
    assert sum(r['closed'] for r in result['activity'])==1
    assert sum(r['started'] for r in result['activity'])==0


def test_combined_filters_and_invalid_options(api_client):
    p=payload(api_client)
    node=api_client.post('/api/nodes',json={'name':'Nodo NOC'}).json()
    person=api_client.post('/api/responsibles',json={'name':'Técnico NOC'}).json()
    api_client.post('/api/internal-tickets',json={**p,'node_id':node['id'],'responsible_id':person['id']})
    api_client.post('/api/internal-tickets',json=p)
    result=summary(api_client,sector_id=p['sector_id'],node_id=node['id'],responsible_id=person['id'],department_id=p['department_id'],incident_type_id=p['incident_type_id'],without_provider=True,status='OPEN')
    assert result['kpis']['started']==1
    assert result['node']['name']=='Nodo NOC'
    for params in ({'timezone':'invalid'}, {'status':'unknown'}, {'provider_id':1,'without_provider':True}, {'from':'2026-01-01T00:00:00Z'}, {'provider_id':0}):
        assert api_client.get(ROOT+'/summary',params=params).status_code==422


def test_statistical_exports_all_details(api_client, db_session):
    from datetime import datetime, timezone
    from ticketyn.models import InternalTicket
    p=payload(api_client)
    db_session.add_all([InternalTicket(**{**p,'start_at':datetime(2026,10,15,14,tzinfo=timezone.utc)},ticket_number=i,reference=f'INT-{i:07}',status='OPEN') for i in range(1,56)])
    db_session.commit()
    assert summary(api_client)['kpis']['started']==55
    response=api_client.get(ROOT+'/export/xlsx')
    assert response.status_code==200,response.text
    book=load_workbook(BytesIO(response.content))
    assert len(list(book['Tickets internos'].values))==56
    assert dict(book['Resumen'].values)['Incidencias iniciadas']==55
    response=api_client.get(ROOT+'/export/pdf')
    assert response.status_code==200 and response.content.startswith(b'%PDF-')
