from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from sqlalchemy import event
from ticketyn.api.fms_exports import capacity_text
from ticketyn.models.fms import Activation, CapacityChange, CircuitLifecycleOperation, ActivationIntervention
from test_fms import operation_payload
from test_ticket_exports import pdf_strings

MODELS = {'activations': Activation, 'capacity-changes': CapacityChange, 'circuit-lifecycle': CircuitLifecycleOperation}


@pytest.fixture(params=list(MODELS))
def exported_operations(request, api_client, operation_payload, ticket_payload, db_session):
    kind = request.param; model = MODELS[kind]
    customer = api_client.post('/api/customers', json={'customer_code': 'EXPORT', 'name': '=Cliente externo'}).json()
    circuit = api_client.post('/api/circuits', json={'customer_id': customer['id'], 'node_id': ticket_payload['node_id'], 'sector_id': ticket_payload['sector_id'], 'circuit_code': 'EXP-C', 'description': '=Circuito externo'}).json()
    reason = api_client.post('/api/deactivation-reasons', json={'name': 'Motivo histórico'}).json()
    records = []
    for index in range(350):
        values = dict(circuit_id=ticket_payload['circuit_id'] if index < 300 else circuit['id'], responsible_id=operation_payload['responsible_id'], effective_at=datetime(2020,1,1 if index % 2 == 0 else 2,12,tzinfo=timezone.utc), sector_id=ticket_payload['sector_id'] if index else None, node_id=ticket_payload['node_id'] if index else None)
        if kind == 'activations':
            values.update(activation_number=index+1, reference=f'ACT-EXP-{index+1:04}')
        elif kind == 'capacity-changes':
            values.update(operation_type='INCREASE' if index%2 == 0 else 'DECREASE', previous_capacity=None if index == 0 else Decimal('100') if index%2 == 0 else Decimal('1000'), previous_unit=None if index == 0 else 'Mbps', new_capacity=Decimal('2000') if index%2 == 0 else Decimal('500'), new_unit='Mbps', description=('Texto extenso & <válido>\n' * 300) if index == 349 else '=Observación')
        else:
            values.update(operation_type='DEACTIVATE' if index%2 == 0 else 'REACTIVATE', reason_id=reason['id'] if index%2 == 0 else None, observations=('Texto extenso & <válido>\n' * 300) if index == 349 else '=Observación')
        records.append(model(**values))
    db_session.add_all(records); db_session.commit()
    if kind == 'activations':
        for content in ('Bitácora primera', 'Bitácora segunda'):
            db_session.add(ActivationIntervention(activation_id=records[0].id, responsible_id=operation_payload['responsible_id'], occurred_at=datetime(2020,1,1,tzinfo=timezone.utc), content=content))
        db_session.commit()
    # Actual Circuit sector/node cannot stand in for historical snapshots.
    node = api_client.post('/api/nodes', json={'name': 'NODO ACTUAL NO HISTÓRICO'}).json()
    sector = api_client.post('/api/sectors', json={'name': 'SECTOR ACTUAL NO HISTÓRICO'}).json()
    assert api_client.patch(f"/api/circuits/{ticket_payload['circuit_id']}", json={'node_id': node['id'], 'sector_id': sector['id']}).status_code == 200
    return kind, customer['id'], circuit['id']


@pytest.mark.parametrize('format', ['pdf', 'xlsx'])
@pytest.mark.parametrize('filters', ['all', 'customer', 'period', 'combined', 'type', 'circuit'])
def test_all_pages_and_exact_list_filters(api_client, exported_operations, format, filters, pdf_strings):
    kind, customer, circuit = exported_operations
    query = {}
    if filters in ('customer', 'combined'): query['customer_id'] = customer
    if filters in ('period', 'combined'): query.update(from_at='2020-01-01T00:00:00Z', to_at='2020-01-02T00:00:00Z')
    if filters == 'type' and kind != 'activations': query['operation_type'] = 'INCREASE' if kind == 'capacity-changes' else 'DEACTIVATE'
    if filters == 'circuit': query['circuit_id'] = circuit
    expected = []
    for offset in (0,200):
        response = api_client.get('/api/' + kind, params={**query, 'limit': 200, 'offset': offset})
        assert response.status_code == 200
        expected += response.json()
    response = api_client.get(f'/api/{kind}/export/{format}', params=query)
    assert response.status_code == 200, response.text
    assert f'ticketyn-{kind}.{format}' in response.headers['content-disposition']
    if format == 'xlsx':
        book = load_workbook(BytesIO(response.content)); sheet = book.worksheets[0]
        assert sheet.max_row == len(expected) + 1
        columns = [c.value for c in sheet[1]]
        assert columns.index('Descripción del circuito') == columns.index('Código del circuito') + 1
        data = [dict(zip(columns, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
        if kind == 'activations': assert [r['Referencia de activación'] for r in data] == [r['reference'] for r in expected]
        assert [r['Código del circuito'] for r in data] == [r['circuit']['circuit_code'] for r in expected]
        assert [r['Fecha y hora efectiva'] for r in data] == [datetime.fromisoformat(r['effective_at']).astimezone(ZoneInfo('America/Guatemala')).isoformat() for r in expected]
        assert [r['Sector'] for r in data] == [r['sector']['name'] if r['sector'] else None for r in expected]
        assert [r['Nodo'] for r in data] == [r['node']['name'] if r['node'] else None for r in expected]
        assert dict(book['Resumen'].values)['Total exportado'] == len(expected)
        for row in sheet.iter_rows(min_row=2):
            for cell in row: assert cell.data_type != 'f'
        if kind == 'capacity-changes' and filters == 'all':
            assert any(r['Capacidad anterior'] == '100 Mbps' and r['Capacidad nueva'] == '2 Gbps' for r in data)
            assert any(r['Capacidad anterior'] == 'Sin registrar' for r in data)
            assert any(r['Valor nuevo registrado (exacto)'] == '2000.000000' and r['Unidad nueva registrada'] == 'Mbps' for r in data)
        if kind == 'circuit-lifecycle': assert all(r['Motivo de baja'] is None for r in data if r['Tipo de operación'] == 'Reactivación')
        book.close()
    else:
        assert response.content.startswith(b'%PDF-') and b'%%EOF' in response.content
        if filters == 'all': Path(f'/tmp/ticketyn-fms-export-{kind}.pdf').write_bytes(response.content)
        content = ' '.join(pdf_strings)
        assert f'Total encontrado: {len(expected)}' in content
        assert 'NODO ACTUAL NO HISTÓRICO' not in content and 'SECTOR ACTUAL NO HISTÓRICO' not in content
        assert any('Página 2 /' in word for word in pdf_strings)
        if kind == 'activations':
            for row in expected: assert row['reference'] in content
            assert 'Bitácora primera' not in content and 'Bitácora segunda' not in content


@pytest.mark.parametrize('kind', list(MODELS))
@pytest.mark.parametrize('format', ['pdf','xlsx'])
def test_empty_and_invalid_exports(api_client, kind, format):
    response = api_client.get(f'/api/{kind}/export/{format}')
    assert response.status_code == 200
    if format == 'xlsx':
        book = load_workbook(BytesIO(response.content)); assert book.worksheets[0].max_row == 1
    else: assert response.content.startswith(b'%PDF')
    for query in [{'customer_id': 0}, {'timezone': 'Unknown/Zone'}, {'from_at': '2020-01-01T00:00:00'}, {'from_at': '2020-01-02T00:00:00Z', 'to_at': '2020-01-01T00:00:00Z'}]:
        assert api_client.get(f'/api/{kind}/export/{format}', params=query).status_code == 422
    if kind != 'activations': assert api_client.get(f'/api/{kind}/export/{format}', params={'operation_type': 'OTHER'}).status_code == 422


@pytest.mark.parametrize('value,unit,expected', [('999','Mbps','999 Mbps'),('1000','Mbps','1 Gbps'),('1758.5','Mbps','1.759 Gbps'),('1.5','Gbps','1.5 Gbps'),('0.5','Gbps','500 Mbps'),('999.9995','Mbps','1000 Mbps'),('100.1235','Mbps','100.124 Mbps')])
def test_capacity_presentation_is_decimal(value, unit, expected):
    original = Decimal(value)
    assert capacity_text(original, unit) == expected
    assert original == Decimal(value)


def test_xlsx_query_count_does_not_grow_with_records(api_client, exported_operations, db_session):
    kind, _, _ = exported_operations
    statements = []
    bind = db_session.get_bind()
    def capture(*args): statements.append(args[2])
    event.listen(bind, 'before_cursor_execute', capture)
    try: assert api_client.get(f'/api/{kind}/export/xlsx').status_code == 200
    finally: event.remove(bind, 'before_cursor_execute', capture)
    # The protected request also validates the principal and security configuration.
    selects = [sql for sql in statements if sql.lstrip().upper().startswith('SELECT ')]
    auth = [sql for sql in selects if 'security_roles' in sql or 'security_config' in sql]
    assert len(auth) == 2
    assert len(selects) - len(auth) == 1


def test_generation_failure_closes_private_temporary_file(api_client, monkeypatch):
    from tempfile import SpooledTemporaryFile
    from ticketyn.api import fms_exports
    files = []
    def temporary(**kwargs):
        file = SpooledTemporaryFile(**kwargs); files.append(file); return file
    def failed(*args): raise RuntimeError('Fallo de generación simulado')
    monkeypatch.setattr(fms_exports, 'SpooledTemporaryFile', temporary)
    monkeypatch.setattr(fms_exports, 'export_pdf', failed)
    with pytest.raises(RuntimeError, match='Fallo de generación simulado'):
        api_client.get('/api/activations/export/pdf')
    assert len(files) == 1 and files[0].closed
