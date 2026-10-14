from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from uuid import uuid4
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine, select, text, inspect
from sqlalchemy.orm import Session
from ticketyn.models import Circuit, Node, Ticket
from ticketyn.models.fms import Activation
from ticketyn.api.fms_data import create_operation
from test_ticket_concurrency import concurrent_database
from test_e2e_update_migration import cli
from test_fms import operation_payload


def node(client, name='Nodo nuevo'):
    response = client.post('/api/nodes', json={'name': name})
    assert response.status_code == 201
    return response.json()['id']


def test_circuit_requires_node_and_legacy_can_be_read(api_client, ticket_payload, db_session):
    body = {'customer_id': ticket_payload['customer_id'], 'sector_id': ticket_payload['sector_id'], 'circuit_code': 'N', 'description': 'Nuevo'}
    assert api_client.post('/api/circuits', json=body).status_code == 422
    assert api_client.post('/api/circuits', json={**body, 'node_id': 99999}).status_code == 404
    row = api_client.post('/api/circuits', json={**body, 'node_id': ticket_payload['node_id']}).json()
    assert row['node_id'] == ticket_payload['node_id']
    circuit = db_session.get(Circuit, row['id']); circuit.node_id = None; db_session.commit()
    assert api_client.get(f"/api/circuits/{row['id']}").json()['node'] is None
    assert api_client.patch(f"/api/circuits/{row['id']}", json={'description': 'Cambio'}).status_code == 422
    assert api_client.patch(f"/api/circuits/{row['id']}", json={'node_id': None}).status_code == 422
    assert api_client.patch(f"/api/circuits/{row['id']}", json={'node_id': ticket_payload['node_id']}).status_code == 200


@pytest.mark.parametrize('path,extra', [
    ('activations', {}),
    ('capacity-changes', {'operation_type': 'INCREASE', 'new_capacity': '500', 'new_unit': 'Mbps'}),
    ('capacity-changes', {'operation_type': 'DECREASE', 'previous_capacity': '1', 'previous_unit': 'Gbps', 'new_capacity': '500', 'new_unit': 'Mbps'}),
    ('circuit-lifecycle', {'operation_type': 'DEACTIVATE'}),
    ('circuit-lifecycle', {'operation_type': 'REACTIVATE'}),
])
def test_fms_assigns_node_and_preserves_history(api_client, ticket_payload, operation_payload, db_session, path, extra):
    extra = dict(extra)
    if extra.get('operation_type') == 'DEACTIVATE':
        extra['reason_id'] = api_client.post('/api/deactivation-reasons', json={'name': 'Prueba'}).json()['id']
    cid = ticket_payload['circuit_id']; circuit = db_session.get(Circuit, cid); circuit.node_id = None; db_session.commit()
    body = {**operation_payload, **extra}
    assert api_client.post('/api/' + path, json=body).status_code == 422
    assert api_client.post('/api/' + path, json={**body, 'node_id': 99999}).status_code == 404
    row = api_client.post('/api/' + path, json={**body, 'node_id': ticket_payload['node_id']})
    assert row.status_code == 201, row.text
    assert row.json()['node_id'] == ticket_payload['node_id']
    assert api_client.get(f'/api/circuits/{cid}').json()['node_id'] == ticket_payload['node_id']
    other = node(api_client)
    assert api_client.post('/api/' + path, json={**body, 'node_id': other}).status_code == 409
    assert api_client.patch(f'/api/circuits/{cid}', json={'node_id': other}).status_code == 200
    assert next(r for r in api_client.get('/api/' + path).json() if r['id'] == row.json()['id'])['node_id'] == ticket_payload['node_id']
    assert db_session.scalar(select(Ticket.id)) is None
    assert api_client.get(f'/api/circuits/{cid}').json()['active']


def test_ticket_snapshots_node_and_filters_exports_stay_historical(api_client, ticket_payload):
    body = {k: v for k,v in ticket_payload.items() if k != 'node_id'}
    first = api_client.post('/api/tickets', json=body).json()
    old = ticket_payload['node_id']; current = node(api_client)
    assert first['node_id'] == old
    assert api_client.patch(f"/api/circuits/{ticket_payload['circuit_id']}", json={'node_id': current}).status_code == 200
    assert api_client.patch(f"/api/tickets/{first['id']}", json={'title': 'Corregido'}).json()['node_id'] == old
    second = api_client.post('/api/tickets', json=body).json()
    assert second['node_id'] == current
    assert api_client.post('/api/tickets', json=ticket_payload).status_code == 409
    third_node = node(api_client, 'No asignado')
    assert api_client.patch(f"/api/tickets/{first['id']}", json={'node_id': third_node}).status_code == 409
    assert [r['id'] for r in api_client.get('/api/tickets', params={'node_id': old}).json()] == [first['id']]
    book = load_workbook(BytesIO(api_client.get('/api/tickets/export/xlsx', params={'node_id': old}).content))
    assert book['Tickets'].max_row == 2
    assert book['Tickets'].cell(2, [c.value for c in book['Tickets'][1]].index('Nodo') + 1).value == 'Nodo de prueba'
    summary = api_client.get('/api/reports/summary', params={'node_id': old}).json()
    assert summary['kpis']['started'] == 1
    other = api_client.post('/api/circuits', json={'customer_id': ticket_payload['customer_id'], 'sector_id': ticket_payload['sector_id'], 'node_id': current, 'circuit_code': 'OTHER', 'description': 'Otro'}).json()
    assert api_client.patch(f"/api/tickets/{first['id']}", json={'circuit_id': other['id']}).json()['node_id'] == current


def test_legacy_ticket_edit_keeps_unknown_history_and_assigns_circuit(api_client, ticket_payload, db_session):
    first = api_client.post('/api/tickets', json=ticket_payload).json()
    cid = ticket_payload['circuit_id']
    ticket = db_session.get(Ticket, first['id']); ticket.node_id = None
    circuit = db_session.get(Circuit, cid); circuit.node_id = None; db_session.commit()
    assert api_client.patch(f"/api/tickets/{first['id']}", json={'title': 'Cambio'}).status_code == 422
    response = api_client.patch(f"/api/tickets/{first['id']}", json={'title': 'Cambio', 'node_id': ticket_payload['node_id']})
    assert response.status_code == 200 and response.json()['node_id'] is None
    assert api_client.get(f'/api/circuits/{cid}').json()['node_id'] == ticket_payload['node_id']


def test_combined_node_sector_assignment_is_atomic(api_client, ticket_payload, operation_payload, db_session):
    cid = ticket_payload['circuit_id']; circuit = db_session.get(Circuit, cid)
    circuit.node_id = circuit.sector_id = None; db_session.commit()
    assert api_client.post('/api/activations', json={**operation_payload, 'node_id': ticket_payload['node_id'], 'sector_id': 99999}).status_code == 404
    saved = api_client.get(f'/api/circuits/{cid}').json()
    assert saved['node_id'] is None and saved['sector_id'] is None
    assert api_client.post('/api/tickets', json={**ticket_payload, 'department_id': 99999}).status_code == 404
    assert api_client.get(f'/api/circuits/{cid}').json()['node_id'] is None
    response = api_client.post('/api/tickets', json=ticket_payload)
    assert response.status_code == 201
    saved = api_client.get(f'/api/circuits/{cid}').json()
    assert saved['node_id'] == ticket_payload['node_id'] and saved['sector_id'] == ticket_payload['sector_id']


def test_concurrent_node_assignment_does_not_overwrite(concurrent_database):
    engine, payload = concurrent_database
    with Session(engine) as session:
        from ticketyn.models import Responsible
        person = Responsible(name='Operador'); other = Node(name='Otro')
        circuit = session.get(Circuit, payload['circuit_id']); circuit.node_id = None
        session.add_all([person, other]); session.commit(); pid, other_id = person.id, other.id
    barrier = Barrier(2)
    def create(nid):
        with Session(engine) as session:
            barrier.wait(timeout=10)
            try:
                return create_operation(session, Activation, {'circuit_id': payload['circuit_id'], 'sector_id': payload['sector_id'], 'node_id': nid, 'responsible_id': pid, 'effective_at': payload['start_at']}).node_id
            except HTTPException as error:
                assert error.status_code == 409
                return None
    with ThreadPoolExecutor(2) as pool:
        results = list(pool.map(create, [payload['node_id'], other_id]))
    assert sum(r is not None for r in results) == 1
    with Session(engine) as session:
        assert session.get(Circuit, payload['circuit_id']).node_id == next(r for r in results if r is not None)
        assert len(list(session.scalars(select(Activation)))) == 1


def test_migration_does_not_invent_nodes_or_change_current(postgres_engine):
    name = 'nodes_' + uuid4().hex
    with postgres_engine.connect().execution_options(isolation_level='AUTOCOMMIT') as admin:
        admin.execute(text('CREATE DATABASE ' + name))
    engine = create_engine(postgres_engine.url.set(database=name))
    try:
        cli(engine.url, 'upgrade', '0016_circuit_sectors')
        with engine.begin() as db:
            db.execute(text("INSERT INTO customers(id,customer_code,name) VALUES(1,'C','Cliente')"))
            db.execute(text("INSERT INTO nodes(id,name) VALUES(1,'Actual')"))
            for table in ('sectors','departments','incident_types'):
                db.execute(text(f"INSERT INTO {table}(id,name) VALUES(1,'Ejemplo')"))
            db.execute(text("INSERT INTO responsibles(id,name) VALUES(1,'Operador')"))
            db.execute(text("INSERT INTO circuits(id,customer_id,circuit_code,description,node_id) VALUES(1,1,'A','Con nodo',1),(2,1,'B','Sin nodo',NULL)"))
            for tid,cid in [(1,1),(2,1),(3,2)]:
                db.execute(text("INSERT INTO tickets(id,ticket_number,reference,title,description,customer_id,circuit_id,sector_id,department_id,incident_type_id,start_at,status) VALUES(:id,:id,:ref,'Histórico','Datos',1,:cid,1,1,1,'2020-01-01T00:00:00Z','OPEN')"), {'id':tid,'cid':cid,'ref':str(tid)})
            db.execute(text("INSERT INTO activations(circuit_id,responsible_id,effective_at,activation_number,reference) VALUES(1,1,'2020-01-01T00:00:00Z',1,'ACT-1')"))
            db.execute(text("INSERT INTO capacity_changes(circuit_id,responsible_id,effective_at,operation_type,new_capacity,new_unit) VALUES(1,1,'2020-01-01T00:00:00Z','INCREASE',500,'Mbps')"))
            db.execute(text("INSERT INTO circuit_lifecycle_operations(circuit_id,responsible_id,effective_at,operation_type,reason_id) VALUES(1,1,'2020-01-01T00:00:00Z','DEACTIVATE',1)"))
            before = db.execute(text('SELECT * FROM tickets ORDER BY id')).all()
            operations_before = {table: db.execute(text(f'SELECT * FROM {table} ORDER BY id')).all() for table in ('activations', 'capacity_changes', 'circuit_lifecycle_operations')}
        cli(engine.url, 'upgrade', 'head')
        with engine.connect() as db:
            assert db.execute(text('SELECT node_id FROM circuits ORDER BY id')).scalars().all() == [1,None]
            rows = db.execute(text('SELECT * FROM tickets ORDER BY id')).all()
            assert [tuple(r[:-1]) for r in rows] == [tuple(r) for r in before]
            assert all(r[-1] is None for r in rows)
            for table, original in operations_before.items():
                restored = db.execute(text(f'SELECT * FROM {table} ORDER BY id')).all()
                assert [tuple(row[:-1]) for row in restored] == [tuple(row) for row in original]
                assert all(row[-1] is None for row in restored)
            for table in ('tickets','activations','capacity_changes','circuit_lifecycle_operations'):
                fk = next(f for f in inspect(db).get_foreign_keys(table) if f['constrained_columns'] == ['node_id'])
                assert fk['referred_table'] == 'nodes' and fk['options'].get('ondelete') != 'CASCADE'
            assert db.scalar(text('SELECT version_num FROM alembic_version')) == '0017_historical_nodes'
    finally:
        engine.dispose()
        with postgres_engine.connect().execution_options(isolation_level='AUTOCOMMIT') as admin:
            admin.execute(text('DROP DATABASE ' + name))
