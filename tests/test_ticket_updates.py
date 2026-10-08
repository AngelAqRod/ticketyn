from datetime import datetime, timezone

import pytest
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from ticketyn.models import Ticket, TicketUpdate


def ticket(api_client, ticket_payload, **changes):
    response = api_client.post('/api/tickets', json={**ticket_payload, **changes})
    assert response.status_code == 201, response.text
    return response.json()


def url(item):
    return f"/api/tickets/{item['id']}/updates"


def create(api_client, item, **changes):
    response = api_client.post(url(item), json={'content': 'Intervención ñ\nVerificación del enlace', **changes})
    assert response.status_code == 201, response.text
    return response.json()


def test_current_defaults_and_server_registration_time(api_client, ticket_payload, db_session):
    item = ticket(api_client, ticket_payload)
    before = datetime.now(timezone.utc)
    update = create(api_client, item)
    after = datetime.now(timezone.utc)
    assert before <= datetime.fromisoformat(update['occurred_at']) <= after
    assert before <= datetime.fromisoformat(update['created_at']) <= after
    assert update['visibility'] == 'INTERNAL'
    assert update['responsible_id'] is None and update['responsible'] is None
    assert update['ticket_id'] == item['id']
    assert api_client.get(url(item)).json() == [update]
    db_session.expire_all()
    parent = db_session.get(Ticket, item['id'])
    # Explicit access to the 1:N association (no implicit eager history on list).
    db_session.refresh(parent, ['updates'])
    assert [entry.id for entry in parent.updates] == [update['id']]


@pytest.mark.parametrize('status', ['OPEN', 'CLOSED'])
def test_retroactive_offset_and_ticket_is_unchanged(api_client, ticket_payload, status):
    item = ticket(api_client, ticket_payload, status=status, end_at='2020-01-01T13:00:00Z')
    before = api_client.get(f"/api/tickets/{item['id']}").json()
    registered_after = datetime.now(timezone.utc)
    update = create(api_client, item, occurred_at='2010-03-15T09:20:31-06:00', visibility='PUBLIC')
    assert datetime.fromisoformat(update['occurred_at']) == datetime(2010, 3, 15, 15, 20, 31, tzinfo=timezone.utc)
    assert datetime.fromisoformat(update['created_at']) >= registered_after
    assert update['occurred_at'] != update['created_at'] and update['visibility'] == 'PUBLIC'
    assert api_client.get(f"/api/tickets/{item['id']}").json() == before


@pytest.mark.parametrize('payload', [
    {}, {'content': ''}, {'content': ' \n\t'}, {'content': None},
    {'content': 'x', 'occurred_at': 'not-a-date'}, {'content': 'x', 'occurred_at': '2026-01-01T12:00:00'},
    {'content': 'x', 'occurred_at': None}, {'content': 'x', 'visibility': 'OTHER'},
    {'content': 'x', 'visibility': None}, {'content': 'x', 'responsible_id': 0},
    {'content': 'x', 'created_at': '2000-01-01T00:00:00Z'}, {'content': 'x', 'ticket_id': 3},
    {'description': 'old field is unsupported'},
])
def test_validation_does_not_insert(api_client, ticket_payload, payload):
    item = ticket(api_client, ticket_payload)
    response = api_client.post(url(item), json=payload)
    assert response.status_code == 422, response.text
    assert api_client.get(url(item)).json() == []


def test_no_created_at_edit_or_delete_endpoint(api_client, ticket_payload):
    item = ticket(api_client, ticket_payload)
    update = create(api_client, item)
    for method in ('patch', 'delete', 'put'):
        response = getattr(api_client, method)(url(item)+f"/{update['id']}")
        assert response.status_code in (404, 405)
    assert api_client.get(url(item)).json() == [update]


def test_responsible_active_missing_and_historical(api_client, ticket_payload):
    item = ticket(api_client, ticket_payload)
    responsible = api_client.post('/api/responsibles', json={'name': 'Operador'}).json()
    update = create(api_client, item, responsible_id=responsible['id'])
    assert update['responsible']['name'] == 'Operador'
    assert api_client.post(url(item), json={'content': 'x', 'responsible_id': 999999}).status_code == 404
    api_client.patch(f"/api/responsibles/{responsible['id']}", json={'active': False})
    assert api_client.post(url(item), json={'content': 'x', 'responsible_id': responsible['id']}).status_code == 422
    stored = api_client.get(url(item)).json()
    assert len(stored) == 1 and stored[0]['responsible']['active'] is False
    assert create(api_client, item, responsible_id=None)['responsible'] is None


def test_ticket_exists_and_path_positive(api_client):
    for method in ('get', 'post'):
        kwargs = {'json': {'content': 'x'}} if method == 'post' else {}
        assert getattr(api_client, method)('/api/tickets/999999/updates', **kwargs).status_code == 404
        assert getattr(api_client, method)('/api/tickets/0/updates', **kwargs).status_code == 422


def test_chronology_uses_instants_and_id_and_isolates_tickets(api_client, ticket_payload):
    first = ticket(api_client, ticket_payload)
    second = ticket(api_client, ticket_payload)
    late = create(api_client, first, occurred_at='2026-10-01T15:00:00Z')
    tied1 = create(api_client, first, occurred_at='2026-09-01T09:00:00-06:00')
    earliest = create(api_client, first, occurred_at='2010-01-01T00:00:00Z')
    tied2 = create(api_client, first, occurred_at='2026-09-01T15:00:00Z')
    create(api_client, second)
    assert [entry['id'] for entry in api_client.get(url(first)).json()] == [earliest['id'], tied1['id'], tied2['id'], late['id']]
    assert len(api_client.get(url(second)).json()) == 1


def test_internal_updates_are_absent_from_existing_exports(api_client, ticket_payload, monkeypatch):
    import io
    from openpyxl import load_workbook
    from reportlab.pdfgen.textobject import PDFTextObject
    strings = []
    original = PDFTextObject._textOut
    def capture(self, value, *args, **kwargs):
        strings.append(value)
        return original(self, value, *args, **kwargs)
    monkeypatch.setattr(PDFTextObject, '_textOut', capture)
    item = ticket(api_client, ticket_payload)
    create(api_client, item, content='INTERNAL-SECRET-MARKER')
    create(api_client, item, content='PUBLIC-MARKER', visibility='PUBLIC')
    # This phase deliberately leaves the export contract unchanged, with no history.
    response = api_client.get('/api/tickets/export/xlsx')
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    values = str(list(workbook['Tickets'].values))
    assert 'INTERNAL-SECRET-MARKER' not in values and 'PUBLIC-MARKER' not in values
    pdf = api_client.get('/api/tickets/export/pdf')
    assert pdf.status_code == 200 and pdf.content.startswith(b'%PDF')
    assert 'INTERNAL-SECRET-MARKER' not in ' '.join(strings)
    report_filters = {'from': '2019-01-01T00:00:00Z', 'to': '2021-01-01T00:00:00Z', 'bucket': 'month', 'tz': 'UTC'}
    response = api_client.get('/api/reports/export/xlsx', params=report_filters)
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert 'INTERNAL-SECRET-MARKER' not in str([list(sheet.values) for sheet in workbook])
    assert api_client.get('/api/reports/export/pdf', params=report_filters).status_code == 200
    assert 'INTERNAL-SECRET-MARKER' not in ' '.join(strings)
    assert 'updates' not in api_client.get(f"/api/tickets/{item['id']}").json()


@pytest.mark.parametrize('values', [{'visibility': 'OTHER'}, {'content': '   '}, {'ticket_id': 999999}, {'responsible_id': 999999}])
def test_db_constraints_reject_invalid_direct_writes(api_client, ticket_payload, db_session, values):
    item = ticket(api_client, ticket_payload)
    with pytest.raises(IntegrityError), db_session.begin_nested():
        db_session.add(TicketUpdate(**{'ticket_id': item['id'], 'content': 'x', **values}))
        db_session.flush()
    assert not list(db_session.scalars(select(TicketUpdate)))


@pytest.mark.parametrize('occurred_at', ['1999-01-01T00:00:00Z', '2030-01-01T12:20:17+09:00'])
def test_any_valid_manual_instant_is_preserved(api_client, ticket_payload, occurred_at):
    item = ticket(api_client, ticket_payload)
    update = create(api_client, item, occurred_at=occurred_at)
    assert datetime.fromisoformat(update['occurred_at']) == datetime.fromisoformat(occurred_at)


def test_history_has_no_responsible_n_plus_one(api_client, ticket_payload, db_session):
    from sqlalchemy import event
    item = ticket(api_client, ticket_payload)
    for index in range(5):
        responsible = api_client.post('/api/responsibles', json={'name': f'Operador {index}'}).json()
        create(api_client, item, responsible_id=responsible['id'])
    statements = []
    connection = db_session.connection()
    def capture(conn, cursor, statement, parameters, context, executemany):
        if statement.lstrip().upper().startswith('SELECT'): statements.append(statement)
    event.listen(connection, 'before_cursor_execute', capture)
    try:
        response = api_client.get(url(item))
        assert response.status_code == 200 and len(response.json()) == 5
        assert len(statements) <= 2
    finally:
        event.remove(connection, 'before_cursor_execute', capture)


def test_cursor_pages_descending_ties_and_concurrent_insert(api_client, ticket_payload):
    item = ticket(api_client, ticket_payload)
    original = [create(api_client, item, occurred_at='2020-01-01T12:00:00.000001Z') for _ in range(23)]
    first = api_client.get(url(item), params={'limit': 10}).json()
    assert [r['id'] for r in first['items']] == [r['id'] for r in reversed(original[-10:])]
    create(api_client, item, occurred_at='2026-01-01T12:00:00Z')
    rows = first['items']
    cursor = first['next_cursor']
    while cursor:
        response = api_client.get(url(item), params={'limit': 10, 'cursor': cursor})
        assert response.status_code == 200
        page = response.json()
        rows.extend(page['items'])
        cursor = page['next_cursor']
    assert [r['id'] for r in rows] == [r['id'] for r in reversed(original)]
    assert len({r['id'] for r in rows}) == 23
    # El consumidor original conserva lista completa y orden ascendente.
    legacy = api_client.get(url(item)).json()
    assert len(legacy) == 24 and legacy[0] == original[0]


@pytest.mark.parametrize('params', [{'limit': 0}, {'limit': 101}, {'limit': 'x'}, {'cursor': 'x'},
                                    {'limit': 10, 'cursor': 'x'}, {'limit': 10, 'cursor': 'e30='}])
def test_invalid_page_parameters(api_client, ticket_payload, params):
    item = ticket(api_client, ticket_payload)
    assert api_client.get(url(item), params=params).status_code == 422


def test_cursor_bound_to_ticket_and_empty_page(api_client, ticket_payload):
    item = ticket(api_client, ticket_payload)
    other = ticket(api_client, ticket_payload)
    assert api_client.get(url(item), params={'limit': 10}).json() == {'items': [], 'next_cursor': None}
    create(api_client, item)
    create(api_client, item)
    cursor = api_client.get(url(item), params={'limit': 1}).json()['next_cursor']
    assert api_client.get(url(other), params={'limit': 10, 'cursor': cursor}).status_code == 422


def test_cursor_preserves_microsecond_order(api_client, ticket_payload):
    item = ticket(api_client, ticket_payload)
    later = create(api_client, item, occurred_at='2020-01-01T00:00:00.000001Z')
    earlier = create(api_client, item, occurred_at='2020-01-01T00:00:00.000000Z')
    first = api_client.get(url(item), params={'limit': 1}).json()
    second = api_client.get(url(item), params={'limit': 1, 'cursor': first['next_cursor']}).json()
    assert first['items'] == [later] and second['items'] == [earlier]
    assert second['next_cursor'] is None
