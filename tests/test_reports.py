from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import event

from test_tickets import create_ticket, RELATED_PATHS

BASE = {'from': '2020-01-01T00:00:00Z', 'to': '2020-01-02T00:00:00Z', 'timezone': 'UTC'}


def summary(client, **params):
    response = client.get('/api/reports/summary', params={**BASE, **params})
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.parametrize('days,grain,length', [(1, 'hour', 24), (7, 'day', 7), (15, 'day', 15), (30, 'day', 30), (150, 'week', 22), (1000, 'month', 33)])
def test_periods_and_zero_buckets(api_client, days, grain, length):
    end = datetime(2020, 1, 1, tzinfo=timezone.utc) + timedelta(days=days)
    report = summary(api_client, to=end.isoformat())
    assert report['period']['granularity'] == grain
    assert len(report['trend']) == length
    assert all(bucket['count'] == 0 for bucket in report['trend'])
    assert report['kpis'] == {'started': 0, 'closed': 0, 'average_duration_seconds': None, 'total_duration_seconds': 0}


def test_temporal_universes_kpis_rankings_and_boundaries(api_client, ticket_payload):
    # Cerrado dentro del período, pero iniciado antes: solo aporta cierres/duración.
    create_ticket(api_client, {**ticket_payload, 'start_at': '2019-12-31T23:00:00Z', 'end_at': '2020-01-01T01:00:00Z', 'status': 'CLOSED'})
    first = create_ticket(api_client, {**ticket_payload, 'start_at': BASE['from'], 'end_at': '2020-01-01T01:00:00Z', 'status': 'CLOSED'})
    create_ticket(api_client, {**ticket_payload, 'end_at': '2020-01-01T13:00:00Z', 'status': 'OPEN'})
    create_ticket(api_client, {**ticket_payload, 'status': 'CLOSED'})  # Sin Fin.
    create_ticket(api_client, {**ticket_payload, 'end_at': BASE['to'], 'status': 'CLOSED'})  # Fin excluido.
    create_ticket(api_client, {**ticket_payload, 'start_at': BASE['to']})  # Inicio excluido.
    report = summary(api_client)
    assert report['kpis'] == {'started': 4, 'closed': 2, 'average_duration_seconds': 5400, 'total_duration_seconds': 10800}
    assert sum(bucket['count'] for bucket in report['trend']) == 4
    assert report['trend'][0]['count'] == 1
    assert report['trend'][12]['count'] == 3
    for kind in ('sectors', 'customers', 'circuits', 'incident_types'):
        assert report[kind][0]['count'] == 4
    assert report['circuits'][0]['customer_id'] == first['customer_id']
    assert report['customers'][0]['label'].startswith('TEST-CUSTOMER')


def test_sector_and_historical_catalogs(api_client, ticket_payload):
    create_ticket(api_client, ticket_payload)
    second_sector = api_client.post('/api/sectors', json={'name': 'Otro sector'}).json()
    create_ticket(api_client, {**ticket_payload, 'sector_id': second_sector['id']})
    for field, path in RELATED_PATHS.items():
        api_client.patch(f'{path}/{ticket_payload[field]}', json={'active': False})
    report = summary(api_client, sector_id=ticket_payload['sector_id'])
    assert report['sector']['name'] == 'Sector de prueba'
    assert report['kpis']['started'] == 1
    assert report['sectors'] == []
    assert report['customers'][0]['label'] == 'TEST-CUSTOMER — Cliente de prueba'
    assert report['circuits'][0]['label'] == 'MANUAL-CIRCUIT'
    assert summary(api_client)['kpis']['started'] == 2
    assert api_client.get('/api/reports/summary', params={**BASE, 'sector_id': 99999}).status_code == 404


@pytest.mark.parametrize('zone,start,end,grain,length', [
    ('Etc/GMT+6', '2020-01-01T00:00:00-06:00', '2020-01-03T00:00:00-06:00', 'day', 2),
    ('America/New_York', '2026-03-08T00:00:00-05:00', '2026-03-09T00:00:00-04:00', 'hour', 23),
    ('America/New_York', '2026-11-01T00:00:00-04:00', '2026-11-02T00:00:00-05:00', 'hour', 25),
])
def test_timezone_buckets_and_dst(api_client, ticket_payload, zone, start, end, grain, length):
    create_ticket(api_client, {**ticket_payload, 'start_at': start})
    report = summary(api_client, **{'from': start, 'to': end, 'timezone': zone, 'granularity': grain})
    assert len(report['trend']) == length
    assert report['trend'][0]['count'] == 1
    assert sum(bucket['count'] for bucket in report['trend']) == 1
    assert len({bucket['start_at'] for bucket in report['trend']}) == length


@pytest.mark.parametrize('changes', [
    {'from': 'invalid'}, {'to': BASE['from']}, {'to': '2019-12-31T00:00:00Z'},
    {'from': '2020-01-01T00:00:00'}, {'timezone': 'Invalid/Zone'}, {'sector_id': 0},
    {'granularity': 'minute'}, {'to': '2035-01-01T00:00:00Z'},
    {'to': '2020-02-01T00:00:00Z', 'granularity': 'hour'},
    {'to': '2022-01-01T00:00:00Z', 'granularity': 'day'},
])
def test_invalid_report_parameters(api_client, changes):
    for path in ('summary', 'export/pdf', 'export/xlsx'):
        assert api_client.get(f'/api/reports/{path}', params={**BASE, **changes}).status_code == 422


def test_all_time_period(api_client):
    response = api_client.get('/api/reports/summary')
    assert response.status_code == 200
    assert response.json()['period']['from_at'] is None
    assert api_client.get('/api/reports/summary', params={'from': BASE['from']}).status_code == 422


@pytest.mark.parametrize('sector_view', [False, True])
def test_xlsx_all_rows_period_sector_styles_and_safe_strings(api_client, ticket_payload, sector_view):
    for _ in range(51):
        create_ticket(api_client, {**ticket_payload, 'title': '=1+1'})
    create_ticket(api_client, {**ticket_payload, 'start_at': BASE['to']})
    other_sector = api_client.post('/api/sectors', json={'name': 'Otro'}).json()['id']
    create_ticket(api_client, {**ticket_payload, 'sector_id': other_sector})
    filters = {**BASE, **({'sector_id': ticket_payload['sector_id']} if sector_view else {})}
    response = api_client.get('/api/reports/export/xlsx', params=filters)
    assert response.status_code == 200, response.text
    assert response.headers['content-type'].startswith('application/vnd.openxmlformats')
    book = load_workbook(BytesIO(response.content))
    expected = 51 if sector_view else 52
    assert book['Tickets'].max_row == expected + 1
    assert ('Sectores' in book.sheetnames) is not sector_view
    metadata = dict(book['Resumen'].values)
    assert metadata['Incidencias iniciadas'] == expected
    assert metadata['Desde (incluido)'] == '2020-01-01T00:00:00+00:00'
    assert book['Tickets']['J2'].data_type == 's'
    assert book['Tickets'].auto_filter.ref == f'A1:S{expected + 1}'
    assert book['Tickets'].freeze_panes == 'A2'
    assert isinstance(book['Tickets']['L2'].value, datetime)
    book.close()


@pytest.mark.parametrize('sector_view', [False, True])
def test_pdf_all_rows_respects_period_sector(api_client, ticket_payload, sector_view, monkeypatch):
    # Captura el texto que ReportLab realmente escribe, sin una dependencia PDF de test.
    from reportlab.pdfgen.canvas import Canvas
    original = Canvas.drawString
    strings = []
    def record(self, x, y, text, *args, **kwargs):
        strings.append(text)
        return original(self, x, y, text, *args, **kwargs)
    monkeypatch.setattr(Canvas, 'drawString', record)
    from reportlab.pdfgen.textobject import PDFTextObject
    original_text = PDFTextObject._textOut
    def record_text(self, text, *args, **kwargs):
        strings.append(text)
        return original_text(self, text, *args, **kwargs)
    monkeypatch.setattr(PDFTextObject, '_textOut', record_text)
    api_client.patch('/api/settings/ticket-number', json={'prefix': 'REP', 'separator': '-'})
    included = [create_ticket(api_client, ticket_payload)['reference'] for _ in range(51)]
    excluded = create_ticket(api_client, {**ticket_payload, 'start_at': BASE['to']})['reference']
    sector = api_client.post('/api/sectors', json={'name': 'Otro'}).json()['id']
    other = create_ticket(api_client, {**ticket_payload, 'sector_id': sector})['reference']
    response = api_client.get('/api/reports/export/pdf', params={**BASE, **({'sector_id': ticket_payload['sector_id']} if sector_view else {})})
    assert response.status_code == 200, response.text
    assert response.content.startswith(b'%PDF-') and b'%%EOF' in response.content
    assert response.headers['content-type'] == 'application/pdf'
    assert all(reference in strings for reference in included)
    assert excluded not in strings
    assert (other in strings) is not sector_view
    assert ('Reporte por Sector' if sector_view else 'Reporte General') in strings
    assert any('Departamentos' in text for text in strings)
    assert any('Tiempo de resolución por departamento' in text for text in strings)


def test_report_query_count_is_constant(api_client, ticket_payload, db_session):
    for _ in range(5):
        create_ticket(api_client, ticket_payload)
    statements = []
    connection = db_session.connection()
    def record(conn, cursor, statement, parameters, context, executemany):
        if statement.lstrip().upper().startswith('SELECT'):
            statements.append(statement)
    event.listen(connection, 'before_cursor_execute', record)
    try:
        assert summary(api_client)['kpis']['started'] == 5
    finally:
        event.remove(connection, 'before_cursor_execute', record)
    assert len(statements) == 19  # 15 consultas existentes + 4 agregaciones de escalamientos; sin N+1.


def test_snapshot_uses_repeatable_read_in_isolated_engine(postgres_engine):
    from sqlalchemy import text
    from sqlalchemy.orm import Session
    from ticketyn.api.reports import read_snapshot
    with Session(postgres_engine) as session:
        read_snapshot(session)
        assert session.scalar(text('SHOW transaction_isolation')) == 'repeatable read'


def test_ranking_full_data_and_deterministic_order(api_client, ticket_payload):
    for index in range(12):
        sector = api_client.post('/api/sectors', json={'name': f'Sector {index:02}'}).json()['id']
        for _ in range(2 if index == 11 else 1):
            create_ticket(api_client, {**ticket_payload, 'sector_id': sector})
    rows = summary(api_client)['sectors']
    assert len(rows) == 12
    assert rows[0]['label'] == 'Sector 11' and rows[0]['count'] == 2
    assert [row['id'] for row in rows[1:]] == sorted(row['id'] for row in rows[1:])
    response = api_client.get('/api/reports/export/xlsx', params=BASE)
    workbook = load_workbook(BytesIO(response.content))
    assert workbook['Sectores'].max_row == 13
    workbook.close()


def test_pdf_and_xlsx_empty_report(api_client):
    for format in ('pdf', 'xlsx'):
        response = api_client.get(f'/api/reports/export/{format}', params=BASE)
        assert response.status_code == 200
        assert 'attachment;' in response.headers['content-disposition']
    assert api_client.get('/api/reports/export/csv', params=BASE).status_code == 422


def test_repeated_hour_counts_and_excel_preserves_offsets(api_client, ticket_payload):
    for start in ('2026-11-01T01:30:00-04:00', '2026-11-01T01:30:00-05:00'):
        create_ticket(api_client, {**ticket_payload, 'start_at': start})
    params = {'from': '2026-11-01T00:00:00-04:00', 'to': '2026-11-02T00:00:00-05:00', 'timezone': 'America/New_York'}
    report = summary(api_client, **params)
    assert [row['count'] for row in report['trend'][:4]] == [0, 1, 1, 0]
    assert report['trend'][1]['start_at'].endswith('-04:00')
    assert report['trend'][2]['start_at'].endswith('-05:00')
    response = api_client.get('/api/reports/export/xlsx', params=params)
    workbook = load_workbook(BytesIO(response.content))
    assert workbook['Tickets']['L2'].value == workbook['Tickets']['L3'].value
    assert {workbook['Tickets']['P2'].value, workbook['Tickets']['P3'].value} == {'2026-11-01T01:30:00-04:00', '2026-11-01T01:30:00-05:00'}
    workbook.close()


def test_export_kpis_match_summary_with_different_start_and_end_universes(api_client, ticket_payload):
    create_ticket(api_client, {**ticket_payload, 'start_at': '2019-12-31T23:00:00Z', 'end_at': '2020-01-01T00:00:00Z', 'status': 'CLOSED'})
    create_ticket(api_client, {**ticket_payload, 'start_at': '2020-01-01T12:00:00Z', 'end_at': '2020-01-01T13:00:00Z', 'status': 'CLOSED'})
    report = summary(api_client)
    response = api_client.get('/api/reports/export/xlsx', params=BASE)
    workbook = load_workbook(BytesIO(response.content))
    metadata = dict(workbook['Resumen'].values)
    assert metadata['Incidencias iniciadas'] == report['kpis']['started'] == 1
    assert metadata['Incidencias cerradas'] == report['kpis']['closed'] == 2
    assert metadata['Duración promedio (segundos)'] == report['kpis']['average_duration_seconds'] == 3600
    assert metadata['Duración acumulada (segundos)'] == report['kpis']['total_duration_seconds'] == 7200
    assert workbook['Tickets'].max_row == 2
    workbook.close()


def test_department_started_and_resolution_universes(api_client, ticket_payload):
    other = api_client.post('/api/departments', json={'name': 'Departamento alterno'}).json()['id']
    create_ticket(api_client, {**ticket_payload, 'start_at': '2019-12-31T23:00:00Z', 'end_at': '2020-01-01T01:00:00Z', 'status': 'CLOSED'})
    create_ticket(api_client, {**ticket_payload, 'start_at': BASE['from'], 'end_at': '2020-01-01T01:00:00Z', 'status': 'CLOSED'})
    create_ticket(api_client, {**ticket_payload, 'status': 'CLOSED'})
    create_ticket(api_client, {**ticket_payload, 'end_at': '2020-01-01T13:00:00Z', 'status': 'OPEN'})
    create_ticket(api_client, {**ticket_payload, 'end_at': BASE['to'], 'status': 'CLOSED'})
    create_ticket(api_client, {**ticket_payload, 'department_id': other, 'start_at': BASE['to']})
    api_client.patch(f"/api/departments/{ticket_payload['department_id']}", json={'active': False})
    report = summary(api_client)
    assert [(row['id'], row['count']) for row in report['departments']] == [(ticket_payload['department_id'], 4)]
    assert report['department_durations'] == [{'id': ticket_payload['department_id'], 'label': 'Departamento de prueba', 'count': 2, 'average_duration_seconds': 5400}]
    assert report['department_durations'] == [{**report['sector_durations'][0], 'id': ticket_payload['department_id'], 'label': 'Departamento de prueba'}]
    assert summary(api_client, sector_id=ticket_payload['sector_id'])['department_durations'] == report['department_durations']
    assert summary(api_client, sector_id=api_client.post('/api/sectors', json={'name': 'Vacío'}).json()['id'])['departments'] == []
    workbook = load_workbook(BytesIO(api_client.get('/api/reports/export/xlsx', params=BASE).content))
    assert workbook['Departamentos']['C2'].value == 4
    assert workbook['Duración por departamento']['C2'].value == 5400
    pdf = api_client.get('/api/reports/export/pdf', params=BASE)
    assert pdf.status_code == 200 and pdf.content.startswith(b'%PDF-')
