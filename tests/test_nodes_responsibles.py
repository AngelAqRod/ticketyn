from io import BytesIO
from pathlib import Path
import runpy
from uuid import uuid4

import pytest
from alembic.migration import MigrationContext
from alembic.operations import Operations
from openpyxl import load_workbook
from sqlalchemy import event, inspect, text

from ticketyn.db.base import Base


@pytest.mark.parametrize('resource', ['nodes', 'responsibles'])
def test_named_catalog_lifecycle(api_client, resource):
    path = f'/api/{resource}'
    response = api_client.post(path, json={'name': 'Centro'})
    assert response.status_code == 201
    item = response.json()
    assert item['active'] and item['created_at']
    assert api_client.get(f'{path}/{item["id"]}').json() == item
    assert api_client.post(path, json={'name': 'Centro'}).status_code == 409
    assert api_client.get(path, params={'search': 'CENT'}).json() == [item]
    inactive = api_client.patch(f'{path}/{item["id"]}', json={'name': 'Nuevo', 'active': False}).json()
    assert api_client.get(path).json() == []
    assert api_client.get(path, params={'include_inactive': True}).json() == [inactive]
    assert api_client.patch(f'{path}/{item["id"]}', json={'active': True}).json()['active']
    assert api_client.get(f'{path}/999999').status_code == 404


@pytest.mark.parametrize('resource', ['nodes', 'responsibles'])
@pytest.mark.parametrize('payload', [{'name': ''}, {'name': '  '}, {'name': None}, {}, {'name': 'X', 'unknown': 1}])
def test_named_catalog_validation(api_client, resource, payload):
    assert api_client.post(f'/api/{resource}', json=payload).status_code == 422


def setup_relations(client, payload):
    node = client.post('/api/nodes', json={'name': 'Nodo Norte'}).json()
    responsible = client.post('/api/responsibles', json={'name': 'Operador Uno'}).json()
    circuit = client.patch(f'/api/circuits/{payload["circuit_id"]}', json={'node_id': node['id']}).json()
    assert circuit['node_id'] == node['id'] and circuit['node']['name'] == node['name']
    return node, responsible


@pytest.mark.parametrize('resource,field,path', [('nodes', 'node_id', 'circuits'), ('responsibles', 'responsible_id', 'tickets')])
@pytest.mark.parametrize('inactive', [False, True])
def test_invalid_assignment(api_client, ticket_payload, resource, field, path, inactive):
    id = 999999
    if inactive:
        id = api_client.post(f'/api/{resource}', json={'name': 'Inactivo', 'active': False}).json()['id']
    if path == 'circuits':
        response = api_client.post('/api/circuits', json={'customer_id': ticket_payload['customer_id'], 'circuit_code': 'NEW', 'description': 'Nuevo', field: id})
    else:
        response = api_client.post('/api/tickets', json={**ticket_payload, field: id})
    assert response.status_code == (422 if inactive else 404)


def test_assign_reassign_unassign_and_history(api_client, ticket_payload):
    node, responsible = setup_relations(api_client, ticket_payload)
    created = api_client.post('/api/tickets', json={**ticket_payload, 'responsible_id': responsible['id']}).json()
    assert created['node']['id'] == node['id']
    assert created['responsible']['id'] == responsible['id']
    for path, item in [('nodes', node), ('responsibles', responsible)]:
        assert api_client.patch(f'/api/{path}/{item["id"]}', json={'active': False}).status_code == 200
    historic = api_client.get(f'/api/tickets/{created["id"]}').json()
    assert not historic['node']['active'] and not historic['responsible']['active']
    assert api_client.patch(f'/api/tickets/{created["id"]}', json={'title': 'Corrección'}).status_code == 200
    replacement = api_client.post('/api/responsibles', json={'name': 'Operador Dos'}).json()
    assert api_client.patch(f'/api/tickets/{created["id"]}', json={'responsible_id': replacement['id']}).json()['responsible']['id'] == replacement['id']
    assert api_client.patch(f'/api/tickets/{created["id"]}', json={'responsible_id': None}).json()['responsible'] is None
    new_node = api_client.post('/api/nodes', json={'name': 'Nodo Dos'}).json()
    assert api_client.patch(f'/api/circuits/{ticket_payload["circuit_id"]}', json={'node_id': new_node['id']}).status_code == 200
    assert api_client.get(f'/api/tickets/{created["id"]}').json()['node']['id'] == new_node['id']
    assert api_client.patch(f'/api/circuits/{ticket_payload["circuit_id"]}', json={'node_id': None}).json()['node'] is None
    assert api_client.get(f'/api/tickets/{created["id"]}').json()['node'] is None


@pytest.mark.parametrize('filters', ['node', 'responsible', 'combined', 'all'])
def test_shared_filters_list_reports_exports(api_client, ticket_payload, filters):
    node, responsible = setup_relations(api_client, ticket_payload)
    created = api_client.post('/api/tickets', json={**ticket_payload, 'responsible_id': responsible['id']}).json()
    api_client.post('/api/tickets', json=ticket_payload)
    query = {}
    if filters in {'node', 'combined', 'all'}: query['node_id'] = node['id']
    if filters in {'responsible', 'combined', 'all'}: query['responsible_id'] = responsible['id']
    if filters == 'all':
        query.update({key: ticket_payload[key] for key in ['customer_id', 'circuit_id', 'sector_id', 'department_id', 'incident_type_id']})
        query.update(status='OPEN', search=ticket_payload['title'], **{'from': '2020-01-01T00:00:00Z', 'to': '2020-01-02T00:00:00Z'})
    rows = api_client.get('/api/tickets', params=query).json()
    assert len(rows) == (2 if filters == 'node' else 1)
    assert rows[0]['node']['id'] == node['id']
    if filters != 'node': assert rows[0]['responsible']['id'] == responsible['id']
    export = api_client.get('/api/tickets/export/xlsx', params=query)
    book = load_workbook(BytesIO(export.content))
    values = list(book['Tickets'].values)
    assert values[0][-2:] == ('Nodo', 'Responsable')
    assert len(values) == len(rows) + 1
    assert values[1][-2] == node['name']
    assert api_client.get('/api/tickets/export/pdf', params=query).content.startswith(b'%PDF')
    period = {'from': '2020-01-01T00:00:00Z', 'to': '2020-01-02T00:00:00Z', 'timezone': 'UTC'}
    period.update({key: value for key, value in query.items() if key in {'node_id', 'responsible_id'}})
    summary = api_client.get('/api/reports/summary', params=period).json()
    assert summary['kpis']['started'] == len(rows)
    for format in ['pdf', 'xlsx']:
        response = api_client.get(f'/api/reports/export/{format}', params=period)
        assert response.status_code == 200
        if format == 'pdf': assert response.content.startswith(b'%PDF')
        else: assert load_workbook(BytesIO(response.content))['Tickets'].max_row == len(rows) + 1


def test_report_new_aggregates_and_nulls(api_client, ticket_payload):
    node, responsible = setup_relations(api_client, ticket_payload)
    api_client.post('/api/tickets', json={**ticket_payload, 'responsible_id': responsible['id'], 'status': 'CLOSED', 'end_at': '2020-01-01T13:00:00Z'})
    # Cierre iniciado fuera del período: aporta cierre y duración, no inicio/ranking.
    api_client.post('/api/tickets', json={**ticket_payload, 'responsible_id': responsible['id'], 'status': 'CLOSED', 'start_at': '2019-12-31T12:00:00Z', 'end_at': '2020-01-01T12:30:00Z'})
    api_client.post('/api/tickets', json={**ticket_payload, 'status': 'CLOSED'})
    for path, item in [('nodes', node), ('responsibles', responsible)]: api_client.patch(f'/api/{path}/{item["id"]}', json={'active': False})
    response = api_client.get('/api/reports/summary', params={'from': '2020-01-01T00:00:00Z', 'to': '2020-01-02T00:00:00Z', 'timezone': 'UTC', 'granularity': 'hour'})
    assert response.status_code == 200
    data = response.json()
    assert data['kpis']['started'] == 2 and data['kpis']['closed'] == 2
    assert data['nodes'][0]['count'] == 2 and data['responsibles'][0]['count'] == 1
    assert len(data['activity']) == len(data['hours']) == 24
    assert sum(row['started'] for row in data['activity']) == 2
    assert sum(row['closed'] for row in data['activity']) == 2
    assert data['hours'][12]['count'] == 2
    assert len(data['weekdays']) == 7 and data['weekdays'][2]['count'] == 2
    assert data['sector_durations'][0]['average_duration_seconds'] == data['kpis']['average_duration_seconds']


def test_circuit_filters_and_null(api_client, ticket_payload):
    node, _ = setup_relations(api_client, ticket_payload)
    query = {'node_id': node['id'], 'customer_id': ticket_payload['customer_id'], 'search': 'principal', 'include_inactive': True}
    # Search using the actual existing description, avoiding fixture-specific assumptions.
    circuit = api_client.get(f'/api/circuits/{ticket_payload["circuit_id"]}').json()
    query['search'] = circuit['description']
    assert api_client.get('/api/circuits', params=query).json() == [circuit]
    api_client.patch(f'/api/circuits/{circuit["id"]}', json={'active': False})
    assert not api_client.get('/api/circuits', params={'node_id': node['id']}).json()
    assert len(api_client.get('/api/circuits', params={**query, 'active': False}).json()) == 1
    plain = api_client.post('/api/circuits', json={'customer_id': ticket_payload['customer_id'], 'circuit_code': 'PLAIN', 'description': 'Sin nodo'}).json()
    assert plain['node_id'] is None and plain['node'] is None


def test_list_query_count_no_n_plus_one(api_client, ticket_payload, db_session):
    node, responsible = setup_relations(api_client, ticket_payload)
    for _ in range(4): api_client.post('/api/tickets', json={**ticket_payload, 'responsible_id': responsible['id']})
    statements = []
    connection = db_session.get_bind()
    listener = lambda *args: statements.append(args[2])
    event.listen(connection, 'before_cursor_execute', listener)
    try:
        assert len(api_client.get('/api/tickets').json()) == 4
    finally:
        event.remove(connection, 'before_cursor_execute', listener)
    assert len(statements) == 1


def test_new_migration_roundtrip(postgres_engine):
    versions = Path(__file__).resolve().parents[1] / 'alembic/versions'
    schema = f'node_migration_{uuid4().hex}'
    with postgres_engine.connect() as connection:
        transaction = connection.begin()
        try:
            connection.execute(text(f'CREATE SCHEMA "{schema}"'))
            connection.execute(text(f'SET LOCAL search_path TO "{schema}"'))
            context = MigrationContext.configure(connection)
            with Operations.context(context):
                for filename in ['0001_initial_catalogs.py', '0002_remove_services.py', '0003_ticket_domain.py']:
                    runpy.run_path(str(versions / filename))['upgrade']()
            connection.execute(text("INSERT INTO customers (customer_code,name) VALUES ('EXISTING','Existing')"))
            connection.execute(text("INSERT INTO circuits (customer_id,circuit_code,description) VALUES (1,'EXISTING.1','Existing')"))
            revision = runpy.run_path(str(versions / '0004_nodes_responsibles.py'))
            assert revision['down_revision'] == '0003_ticket_domain'
            with Operations.context(context): revision['upgrade']()
            inspector = inspect(connection)
            # This historical roundtrip stops at 0004, before follow-up in 0008.
            assert set(inspector.get_table_names(schema=schema)) == set(Base.metadata.tables) - {'ticket_updates', 'escalation_reasons', 'ticket_escalations', 'positions'}
            assert connection.scalar(text('SELECT node_id FROM circuits')) is None
            for table, column, target in [('circuits','node_id','nodes'),('tickets','responsible_id','responsibles')]:
                assert next(c for c in inspector.get_columns(table, schema=schema) if c['name'] == column)['nullable']
                assert any(fk['constrained_columns'] == [column] and fk['referred_table'] == target for fk in inspector.get_foreign_keys(table,schema=schema))
                assert any(index['name'] == f'ix_{table}_{column}' for index in inspector.get_indexes(table,schema=schema))
            assert 'node_id' not in Base.metadata.tables['tickets'].c
            assert 'responsible_id' not in Base.metadata.tables['circuits'].c
            assert 'responsible_id' not in Base.metadata.tables['nodes'].c
            assert not any('hub' in name for name in Base.metadata.tables)
            for table in ['nodes', 'responsibles']:
                assert connection.scalar(text(f'SELECT count(*) FROM {table}')) == 0
            with Operations.context(context): revision['downgrade']()
            assert 'nodes' not in inspect(connection).get_table_names(schema=schema)
            assert 'responsibles' not in inspect(connection).get_table_names(schema=schema)
            assert connection.scalar(text('SELECT circuit_code FROM circuits')) == 'EXISTING.1'
            assert len(inspect(connection).get_foreign_keys('tickets', schema=schema)) == 5
        finally:
            transaction.rollback()

@pytest.mark.parametrize('field,value', [('responsible_id', -1), ('responsible_id', 0), ('node_id', 1)])
def test_ticket_assignment_input_contract(api_client, ticket_payload, field, value):
    assert api_client.post('/api/tickets', json={**ticket_payload, field: value}).status_code == 422


@pytest.mark.parametrize('mode', ['node_id', 'responsible_id'])
def test_report_assignment_parameters(api_client, mode):
    period = {'from': '2020-01-01T00:00:00Z', 'to': '2020-01-02T00:00:00Z'}
    assert api_client.get('/api/reports/summary', params={**period, mode: 0}).status_code == 422
    assert api_client.get('/api/reports/summary', params={**period, mode: 999999}).status_code == 404


def test_exports_complete_and_readable(api_client, ticket_payload):
    node, responsible = setup_relations(api_client, ticket_payload)
    for _ in range(51):
        response = api_client.post('/api/tickets', json={**ticket_payload, 'responsible_id': responsible['id']})
        assert response.status_code == 201
    filters = {'node_id': node['id'], 'responsible_id': responsible['id']}
    period = {**filters, 'from': '2020-01-01T00:00:00Z', 'to': '2020-01-02T00:00:00Z', 'timezone': 'UTC'}
    general = api_client.get('/api/reports/export/pdf', params={key: value for key, value in period.items() if key not in filters})
    assert general.status_code == 200
    Path('/tmp/ticketyn-report-general-composition.pdf').write_bytes(general.content)
    for path, params in [('tickets', filters), ('reports', period)]:
        for format in ['pdf', 'xlsx']:
            response = api_client.get(f'/api/{path}/export/{format}', params=params)
            assert response.status_code == 200
            if format == 'xlsx':
                book = load_workbook(BytesIO(response.content))
                assert book['Tickets'].max_row == 52
                assert 'Nodo' in [cell.value for cell in book['Tickets'][1]]
                assert 'Responsable' in [cell.value for cell in book['Tickets'][1]]
            else:
                assert response.content.startswith(b'%PDF') and response.content.rstrip().endswith(b'%%EOF')
            # Artifacts only from the isolated test database, for PDF-reader verification.
            Path(f'/tmp/ticketyn-{path}-node-responsible.{format}').write_bytes(response.content)


def test_ticket_pdf_starts_table_on_first_page(api_client, ticket_payload, monkeypatch):
    from reportlab.pdfgen.textobject import PDFTextObject
    node, responsible = setup_relations(api_client, ticket_payload)
    references = set()
    for _ in range(25):
        references.add(api_client.post('/api/tickets', json={**ticket_payload, 'responsible_id': responsible['id']}).json()['reference'])
    pages = []
    original = PDFTextObject._textOut
    def capture(self, value, *args, **kwargs):
        if value in references:
            pages.append(self._canvas._pageNumber)
        return original(self, value, *args, **kwargs)
    monkeypatch.setattr(PDFTextObject, '_textOut', capture)
    response = api_client.get('/api/tickets/export/pdf')
    assert response.status_code == 200
    assert 1 in pages


@pytest.mark.parametrize('timezone, hour, weekday', [('UTC', 1, 2), ('Etc/GMT+6', 19, 1)])
def test_new_distributions_respect_timezone(api_client, ticket_payload, timezone, hour, weekday):
    node, responsible = setup_relations(api_client, ticket_payload)
    api_client.post('/api/tickets', json={**ticket_payload, 'responsible_id': responsible['id'], 'status': 'CLOSED', 'start_at': '2020-01-01T01:00:00Z', 'end_at': '2020-01-01T02:00:00Z'})
    query = {'from': '2019-12-31T00:00:00Z', 'to': '2020-01-02T00:00:00Z', 'timezone': timezone, 'node_id': node['id'], 'responsible_id': responsible['id']}
    response = api_client.get('/api/reports/summary', params=query)
    assert response.status_code == 200
    data = response.json()
    assert data['hours'][hour]['count'] == 1
    assert data['weekdays'][weekday]['count'] == 1
    assert sum(row['started'] for row in data['activity']) == 1
    assert sum(row['closed'] for row in data['activity']) == 1
    assert data['kpis']['total_duration_seconds'] == 3600


def test_unassigned_exports_and_pdf_friendly_names(api_client, ticket_payload, monkeypatch):
    from reportlab.pdfgen.textobject import PDFTextObject
    api_client.post('/api/tickets', json=ticket_payload)
    response = api_client.get('/api/tickets/export/xlsx')
    rows = list(load_workbook(BytesIO(response.content))['Tickets'].values)
    assert rows[1][-2:] == (None, None)
    node, responsible = setup_relations(api_client, ticket_payload)
    api_client.post('/api/tickets', json={**ticket_payload, 'responsible_id': responsible['id']})
    strings = []
    original = PDFTextObject._textOut
    def capture(self, value, *args, **kwargs):
        strings.append(value)
        return original(self, value, *args, **kwargs)
    monkeypatch.setattr(PDFTextObject, '_textOut', capture)
    response = api_client.get('/api/tickets/export/pdf', params={'node_id': node['id'], 'responsible_id': responsible['id']})
    assert response.status_code == 200
    content = ' '.join(strings)
    assert 'Nodo: Nodo Norte' in content and 'Responsable: Operador Uno' in content
    assert 'node_id=' not in content and 'responsible_id=' not in content
