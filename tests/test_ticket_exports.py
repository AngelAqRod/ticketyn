from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import event

from test_tickets import create_ticket, RELATED_PATHS


def workbook(client, **params):
    response = client.get('/api/tickets/export/xlsx', params=params)
    assert response.status_code == 200, response.text
    return load_workbook(BytesIO(response.content))


@pytest.fixture
def pdf_strings(monkeypatch):
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.pdfgen.textobject import PDFTextObject
    strings = []
    original, original_text = Canvas.drawString, PDFTextObject._textOut
    def draw(self, x, y, text, *args, **kwargs):
        strings.append(text)
        return original(self, x, y, text, *args, **kwargs)
    def text_out(self, text, *args, **kwargs):
        strings.append(text)
        return original_text(self, text, *args, **kwargs)
    monkeypatch.setattr(Canvas, 'drawString', draw)
    original_right = Canvas.drawRightString
    def draw_right(self, x, y, text, *args, **kwargs):
        strings.append(text)
        return original_right(self, x, y, text, *args, **kwargs)
    monkeypatch.setattr(Canvas, 'drawRightString', draw_right)
    monkeypatch.setattr(PDFTextObject, '_textOut', text_out)
    return strings


@pytest.mark.parametrize('format', ['pdf', 'xlsx'])
def test_empty_exports_are_valid(api_client, format):
    response = api_client.get(f'/api/tickets/export/{format}')
    assert response.status_code == 200
    assert 'attachment;' in response.headers['content-disposition']
    if format == 'pdf':
        assert response.content.startswith(b'%PDF-') and b'%%EOF' in response.content
    else:
        book = load_workbook(BytesIO(response.content))
        assert book.sheetnames == ['Tickets']
        assert book['Tickets'].max_row == 1
        book.close()


def test_xlsx_complete_historical_data_and_headers(api_client, ticket_payload):
    created = create_ticket(api_client, {**ticket_payload, 'title': '=1+1 <&> ñ', 'description': 'Descripción con acentos & < >', 'end_at': '2020-01-01T13:00:00Z'})
    for field, path in RELATED_PATHS.items():
        api_client.patch(f'{path}/{ticket_payload[field]}', json={'active': False})
    book = workbook(api_client, timezone='Etc/GMT+6')
    sheet = book['Tickets']
    assert sheet.freeze_panes == 'A2' and sheet.auto_filter.ref == 'A1:Q2'
    assert [cell.value for cell in sheet[1]][:6] == ['Referencia', 'Estado', 'Cliente', 'Código cliente', 'Circuito', 'Código circuito']
    assert sheet['A2'].value == created['reference']
    assert sheet['J2'].value == '=1+1 <&> ñ' and sheet['J2'].data_type == 's'
    assert sheet['K2'].value == 'Descripción con acentos & < >'
    assert sheet['L2'].value.endswith('-06:00')
    assert sheet['N2'].value == '1 h' and sheet['O2'].value == 3600
    assert sheet['C2'].value == 'Cliente de prueba'
    book.close()


@pytest.mark.parametrize('field', ['status', 'customer_id', 'circuit_id', 'sector_id', 'department_id', 'incident_type_id', 'search', 'from', 'to'])
def test_every_filter_matches_list_semantics(api_client, ticket_payload, field):
    wanted = create_ticket(api_client, {**ticket_payload, 'title': 'Único objetivo', 'start_at': '2020-01-01T12:00:00Z'})
    other = dict(ticket_payload, title='Segundo registro', start_at='2020-01-02T12:00:00Z', status='CLOSED')
    if field in ('customer_id', 'circuit_id'):
        customer = api_client.post('/api/customers', json={'customer_code': 'OTHER', 'name': 'Otro'}).json()['id']
        circuit = api_client.post('/api/circuits', json={'customer_id': customer, 'circuit_code': 'OTHER.C', 'description': 'Otro'}).json()['id']
        other.update(customer_id=customer, circuit_id=circuit)
    elif field in RELATED_PATHS:
        other[field] = api_client.post(RELATED_PATHS[field], json={'name': 'Otro'}).json()['id']
    created_other = create_ticket(api_client, other)
    value = wanted.get(field)
    if field == 'search': value = 'objetivo'
    if field == 'from': value = '2020-01-02T00:00:00Z'
    if field == 'to': value = '2020-01-02T00:00:00Z'
    params = {field: value}
    expected = api_client.get('/api/tickets', params=params).json()
    book = workbook(api_client, **params)
    assert {row[0] for row in list(book['Tickets'].values)[1:]} == {row['reference'] for row in expected}
    assert len(expected) == 1
    if field == 'from': assert expected[0]['reference'] == created_other['reference']
    book.close()


@pytest.mark.parametrize('format', ['pdf', 'xlsx'])
def test_all_results_ignoring_pagination_and_combination(api_client, ticket_payload, format, pdf_strings):
    refs = [create_ticket(api_client, {**ticket_payload, 'title': 'Coincide'})['reference'] for _ in range(53)]
    excluded = create_ticket(api_client, {**ticket_payload, 'status': 'CLOSED'})['reference']
    params = {**{field: ticket_payload[field] for field in RELATED_PATHS}, 'status': 'OPEN', 'search': 'Coincide',
        'from': '2020-01-01T00:00:00Z', 'to': '2020-01-02T00:00:00Z', 'limit': 1, 'offset': 50}
    response = api_client.get(f'/api/tickets/export/{format}', params=params)
    assert response.status_code == 200, response.text
    if format == 'xlsx':
        book = load_workbook(BytesIO(response.content))
        assert book['Tickets'].max_row == 54
        assert {row[0] for row in list(book['Tickets'].values)[1:]} == set(refs)
        book.close()
    else:
        assert all(ref in pdf_strings for ref in refs)
        assert excluded not in pdf_strings
        assert any('Página 2 /' in value for value in pdf_strings)
        assert 'Referencia' in pdf_strings


def test_pdf_friendly_context_special_text_and_inactive_history(api_client, ticket_payload, pdf_strings):
    created = create_ticket(api_client, {**ticket_payload, 'title': 'Revisión <&> ÁÉ ñ'})
    for field, path in RELATED_PATHS.items(): api_client.patch(f'{path}/{ticket_payload[field]}', json={'active': False})
    response = api_client.get('/api/tickets/export/pdf', params={'customer_id': ticket_payload['customer_id'], 'status': 'OPEN', 'timezone': 'Etc/GMT+6'})
    assert response.status_code == 200
    assert created['reference'] in pdf_strings
    text = ' '.join(pdf_strings)
    assert 'Listado de tickets' in text and 'Cliente de prueba' in text
    assert 'Estado: Abierto' in text and 'customer_id=' not in text
    assert 'Revisión' in text and '<' in text and 'ñ' in text
    assert 'Total encontrado: 1' in text


@pytest.mark.parametrize('params', [{'timezone': 'Bad/Zone'}, {'from': '2020-01-01T00:00:00'}, {'from': '2020-02-01T00:00:00Z', 'to': '2020-01-01T00:00:00Z'}, {'customer_id': 0}, {'status': 'UNKNOWN'}])
def test_invalid_export_parameters(api_client, params):
    for format in ('pdf', 'xlsx'):
        assert api_client.get(f'/api/tickets/export/{format}', params=params).status_code == 422


def test_export_query_count_is_independent_of_rows(api_client, ticket_payload, db_session):
    for _ in range(8): create_ticket(api_client, ticket_payload)
    queries = []
    connection = db_session.connection()
    def record(conn, cursor, statement, parameters, context, many):
        if statement.lstrip().upper().startswith('SELECT'): queries.append(statement)
    event.listen(connection, 'before_cursor_execute', record)
    try:
        assert api_client.get('/api/tickets/export/xlsx').status_code == 200
    finally: event.remove(connection, 'before_cursor_execute', record)
    assert len(queries) == 1
